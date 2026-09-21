# -*- coding: utf-8 -*-
# ══════════════════════════════════════════════════════════════
# ⚙️  KORTEK Backend v1.0
#    자동화 메일 · Telegram 알람 · AI 분석 서버
#    포트: 5000 단일 운영
#
#    [포함 기능]
#    /health            — 서버 상태 확인
#    /config            — SMTP 설정 저장
#    /send-mail         — 메일 발송 (SMTP)
#    /fetch-mail        — 메일 수신 (POP3)
#    /send-telegram     — Telegram 메시지 발송
#    /telegram/config   — Telegram 설정 저장/조회
#    /telegram/members  — 팀원 관리 (등록/조회/삭제)
#    /telegram/test     — Telegram 발송 테스트
#    /telegram/encrypt  — Telegram 설정 암호화 (Drive 저장용)
#    /telegram/decrypt  — Telegram 설정 복호화 (Drive 로드용)
#    /mail/encrypt      — SMTP 설정 암호화 (Drive 저장용)
#    /mail/decrypt      — SMTP 설정 복호화 (Drive 로드용)
#    /all/encrypt       — 전체 설정 암호화 (Drive 저장용)
#    /all/decrypt       — 전체 설정 복호화 (Drive 로드용)
#    /schedule          — 예약 발송 규칙 등록/조회 (GET/POST)
#    /schedule/<id>     — 예약 발송 규칙 삭제 (DELETE)
#    /sap-fetch         — [2026-09-14 신규] AI 문답 "SAP 조회" — 이미 로그인된 SAP GUI의
#                          현재 화면을 텍스트로 덤프 (SAP GUI Scripting, 비밀번호 저장 없음)
#
#    [2026-08-31 신규] 예약 발송(반복 규칙) 스케줄러
#    기존엔 "언제 보낼지" 판단을 전부 브라우저(JS setInterval)가 맡고 있어서,
#    브라우저 탭이 열려 있어야만 알람/공지가 발송됐다. 이제 브라우저는 "무엇을
#    보낼지"(규칙)만 이 서버에 등록하고, "언제 보낼지"는 서버가 1분마다 자체
#    확인해서 직접 발송한다 — 탭이 꺼져 있어도 이 서버(kortek_backend.bat)만
#    켜져 있으면 계속 동작한다. 서버가 꺼져있던 동안 지나간 시각은 그냥
#    건너뛴다(밀린 발송을 몰아서 보내지 않음).
# ══════════════════════════════════════════════════════════════

import sys
import os, json, re, poplib, email, smtplib, hashlib, base64, html, threading, time, uuid, subprocess

# 💡 [2026-09-14] sap_bridge_32.py에서 겪은 것과 같은 부류의 문제(콘솔이 실제 콘솔이 아니라
#    파이프/다른 인코딩으로 연결되면 Windows에서 stdout이 cp949로 잡혀 한글 print()가
#    UnicodeEncodeError로 죽을 수 있음)를 이 파일도 예방 차원에서 방어. kortek_backend.bat의
#    `chcp 65001`은 실제 콘솔에 붙어 있을 때만 효과가 있어, 다른 실행 방식(서비스 등록,
#    다른 셸 등)에서도 안전하도록 명시적으로 재설정.
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')
from email.header     import decode_header
from email.utils      import parsedate_to_datetime
from email.mime.text  import MIMEText
from email.mime.multipart import MIMEMultipart
from datetime         import datetime, timezone, timedelta
from urllib.parse     import quote

import requests
from flask      import Flask, request, jsonify, g
from flask_cors import CORS

try:
    from cryptography.fernet import Fernet, InvalidToken
    CRYPTO_OK = True
except ImportError:
    CRYPTO_OK = False
    print("[경고] cryptography 미설치 → pip install cryptography")

try:
    from google.oauth2 import service_account as _google_service_account
    from google.auth.transport.requests import Request as _GoogleAuthRequest
    GOOGLE_AUTH_OK = True
except ImportError:
    GOOGLE_AUTH_OK = False
    print("[경고] google-auth 미설치 → pip install google-auth (업무별 예약 알람의 실시간 조회 기능에 필요)")

app = Flask(__name__)
# 💡 [2026-09-14] SAP GUI Scripting(win32com)은 이 파일(64비트 Python)에서 직접 import하지 않는다
#    — 아래 /sap-fetch 라우트 바로 위 주석 참고. 32비트 전용 코드/의존성은 전부 sap_bridge_32.py 쪽.

# ── CORS: 허용 출처 제한 (보안 강화) ─────────────────────────
ALLOWED_ORIGINS = [
    "https://yhparkkortek.github.io",
    "http://127.0.0.1:5500",
    "http://localhost:5500",
    "null",   # 로컬 HTML 파일 (file://)
]

@app.after_request
def add_cors_headers(resp):
    origin = request.headers.get("Origin", "")
    if origin in ALLOWED_ORIGINS or not origin:
        resp.headers["Access-Control-Allow-Origin"]  = origin or "*"
        resp.headers["Access-Control-Allow-Methods"] = "GET, POST, DELETE, OPTIONS"
        resp.headers["Access-Control-Allow-Headers"] = "Content-Type"
    return resp

# ── 공통 상수 ─────────────────────────────────────────────────
KST            = timezone(timedelta(hours=9))
BASE_DIR       = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE    = os.path.join(BASE_DIR, 'mail_config.json')
TG_CONFIG_FILE = os.path.join(BASE_DIR, 'telegram_config.json')
SCHEDULE_FILE  = os.path.join(BASE_DIR, 'schedule_rules.json')
POP3_HOST      = "gw.kortek.co.kr"
POP3_PORT      = 110

# ── 업무별(개별 태스크) 예약 알람 — 구글드라이브 서비스계정 읽기전용 접근 ──
#    GANTT_CHART_V02_Color.html(js/04a-core-app-globals.js의 SHARED_FOLDER_ID)과 동일한 폴더.
#    이 폴더 밑에 프로젝트 파일들과 App_Config/AddressBook_Shared.json이 들어있다.
GANTT_SHARED_FOLDER_ID     = '1ldb3Bc7dNNSKKgmNviw43aCgrvxQG9bS'
GOOGLE_SERVICE_ACCOUNT_FILE = os.environ.get(
    'GOOGLE_SERVICE_ACCOUNT_FILE',
    os.path.join(BASE_DIR, 'google_service_account.json')
)
ADDRESS_BOOK_CACHE_TTL_SEC = 300  # 주소록은 자주 안 바뀌므로 5분 캐시(매 발송 tick마다 Drive 조회하지 않도록)

# ── SMTP 기본값 ───────────────────────────────────────────────
SMTP_HOST = ''
SMTP_PORT = 25
SMTP_USER = ''
SMTP_PASS = ''
SMTP_TLS  = True

# ── Telegram 런타임 상태 (메모리에만 유지) ────────────────────
TELEGRAM_TOKEN   = ''
TELEGRAM_CHAT_ID = ''
TELEGRAM_MEMBERS = []

# ── 예약 발송 규칙 (메모리 + schedule_rules.json 영속 저장) ────
#    스케줄러 백그라운드 스레드와 Flask 요청 스레드가 동시에 건드릴 수 있어 락으로 보호.
SCHEDULE_RULES = []
_SCHEDULE_LOCK = threading.Lock()


# ══════════════════════════════════════════════════════════════
# ⚙️ 설정 로드
# ══════════════════════════════════════════════════════════════
def load_smtp_config():
    global SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASS
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, 'r', encoding='utf-8') as f:
                cfg = json.load(f)
            SMTP_HOST = cfg.get('host', SMTP_HOST)
            SMTP_PORT = int(cfg.get('port', SMTP_PORT))
            SMTP_USER = cfg.get('user', SMTP_USER)
            SMTP_PASS = cfg.get('pass', SMTP_PASS)
        except Exception as e:
            print(f"[경고] mail_config.json 로드 실패: {e}")

def load_tg_config() -> dict:
    if os.path.exists(TG_CONFIG_FILE):
        try:
            with open(TG_CONFIG_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception as e:
            print(f"[경고] telegram_config.json 로드 실패: {e}")
    return {"token": "", "default_chat_id": "", "members": []}

def save_tg_config(cfg: dict):
    with open(TG_CONFIG_FILE, 'w', encoding='utf-8') as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)

def load_schedule_rules() -> list:
    if os.path.exists(SCHEDULE_FILE):
        try:
            with open(SCHEDULE_FILE, 'r', encoding='utf-8') as f:
                data = json.load(f)
                return data if isinstance(data, list) else []
        except Exception as e:
            print(f"[경고] schedule_rules.json 로드 실패: {e}")
    return []

def save_schedule_rules(rules: list):
    try:
        with open(SCHEDULE_FILE, 'w', encoding='utf-8') as f:
            json.dump(rules, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"[경고] schedule_rules.json 저장 실패: {e}")

# 서버 시작 시 로드
load_smtp_config()
_tg = load_tg_config()
TELEGRAM_TOKEN   = _tg.get("token", "")
TELEGRAM_CHAT_ID = _tg.get("default_chat_id", "")
TELEGRAM_MEMBERS = _tg.get("members", [])
SCHEDULE_RULES   = load_schedule_rules()


# ══════════════════════════════════════════════════════════════
# 🔐 암호화 유틸 (Fernet + SHA-256)
# ══════════════════════════════════════════════════════════════
def password_to_key(password: str) -> bytes:
    """비밀번호 문자열 → Fernet 키 (SHA-256 해시 후 Base64)"""
    hashed = hashlib.sha256(password.encode('utf-8')).digest()
    return base64.urlsafe_b64encode(hashed)

def encrypt_data(data: dict, password: str) -> str:
    """dict → 암호화 문자열"""
    if not CRYPTO_OK:
        raise RuntimeError("cryptography 라이브러리 미설치")
    key     = password_to_key(password)
    f       = Fernet(key)
    payload = json.dumps(data, ensure_ascii=False).encode('utf-8')
    return f.encrypt(payload).decode('utf-8')

def decrypt_data(encrypted: str, password: str) -> dict:
    """암호화 문자열 → dict"""
    if not CRYPTO_OK:
        raise RuntimeError("cryptography 라이브러리 미설치")
    key  = password_to_key(password)
    f    = Fernet(key)
    raw  = f.decrypt(encrypted.encode('utf-8'))
    return json.loads(raw.decode('utf-8'))


# ══════════════════════════════════════════════════════════════
# 📨 메일 유틸
# ══════════════════════════════════════════════════════════════
def decode_str(value):
    if not value:
        return ""
    parts, result = decode_header(value), []
    for part, enc in parts:
        if isinstance(part, bytes):
            result.append(part.decode(enc or "utf-8", errors="replace"))
        else:
            result.append(part)
    return "".join(result)

def extract_body(msg):
    body = ""
    plain_fallback = ""  # 💡 text/html 파트가 끝내 없는 텍스트 전용 메일 대비 폴백
    # 💡 [2026-08-28 버그 수정] 표(HTML table)가 있는 메일에서 "경계문자(│)로 칸을 구분해주기로 했는데도
    #    표가 여전히 알아볼 수 없게 나온다"는 지적 — 원인은 아래 우선순위였다. multipart/alternative
    #    메일은 보통 text/plain 파트도 같이 들어있는데, 그건 Outlook 등 메일 클라이언트가 "표를 이미
    #    한 줄씩 늘어놓은 형태로 미리 납작하게 만들어둔" 버전이라 칸 구분 정보가 아예 없다. 그런데 예전
    #    코드는 그 text/plain을 무조건 최우선으로 골라 썼고, 표 구조가 살아있는 text/html은 text/plain이
    #    "아예 없을 때"만 쓰는 최후 폴백으로 밀려나 있었다 — 그래서 아래 CELL_BOUNDARY 처리(표 칸마다
    #    │ 경계문자를 넣어주는 로직)가 있어도 실행될 기회 자체가 없었다. text/html을 우선으로 바꾸고,
    #    text/html이 없는 경우에만 text/plain으로 폴백하도록 순서를 뒤집는다.
    if msg.is_multipart():
        for part in msg.walk():
            ct = part.get_content_type()
            cd = str(part.get("Content-Disposition", ""))
            if "attachment" in cd:
                continue
            if ct == "text/html" and not body:
                charset = part.get_content_charset() or "utf-8"
                try:
                    body = part.get_payload(decode=True).decode(charset, errors="replace")
                except Exception:
                    pass
            elif ct == "text/plain" and not plain_fallback:
                charset = part.get_content_charset() or "utf-8"
                try:
                    plain_fallback = part.get_payload(decode=True).decode(charset, errors="replace")
                except Exception:
                    pass
        if not body and plain_fallback:
            body = plain_fallback
    else:
        charset = msg.get_content_charset() or "utf-8"
        try:
            body = msg.get_payload(decode=True).decode(charset, errors="replace")
        except Exception:
            body = ""
    # 💡 [2026-08-25 버그 수정] "원문 보기"에서 메일 속 표(HTML table)가 다 깨져 보인다는 지적 — 원인은
    #    두 가지였다. ① <table>/<tr>/<td>/<th> 태그를 아무 구분자 없이 그냥 다 지워버려서, 표 칸
    #    내용이 "Item1Qty1Price1"처럼 서로 다 붙어버렸음(줄바꿈 성격인 <br>/<p>/<div>만 개행으로
    #    바꿔주고 있었음). ② 그 뒤 공백·탭을 무조건 한 칸으로 압축해서, 설령 칸이 스페이스로 정렬된
    #    표(예: 고정폭 폰트 기준 정렬)라 해도 그 정렬용 공백까지 다 뭉개졌음.
    #    → 표 태그도 <br>/<p>/<div>처럼 개행/탭으로 먼저 바꿔 구조를 보존하고(①), 공백 압축은
    #    "짧은 공백(표 정렬용)"은 그대로 두고 "지나치게 긴 공백(HTML 레이아웃 찌꺼기)"만 적당히 줄이도록
    #    완화했다(②).
    # 💡 [2026-08-26 버그 수정 ③] ①②만으로는 부족했다 — 실제 메일(특히 Outlook)의 표 HTML은 보통
    #    예쁘게 들여쓰기되어 있어서 <td>...</td> 사이에 "줄바꿈+공백"만으로 된 서식용 텍스트가 이미
    #    원본에 끼어있다. 우리가 넣는 탭(\t)이 하필 그 원본 줄바꿈 바로 앞에 붙게 되고, 아래
    #    ln.rstrip()이 "줄 끝 찌꺼기"로 착각해 탭까지 같이 지워버려서 — 칸 구분자가 사라지고 칸마다
    #    다시 다른 줄로 떨어졌다. 태그 사이에 순수 공백/줄바꿈만 있는 서식용 여백을 먼저 다 걷어내서,
    #    최종적으로 남는 탭·줄바꿈은 전부 우리가 의도적으로 넣은 것만 남게 한다.
    # 💡 [2026-08-26 변경] 칸 구분자를 탭(\t) 대신 눈에 보이는 경계문자 " │ "(U+2502)로 변경 — 탭은
    #    칸 값 길이가 다르면 세로줄이 안 맞아 보이는데, │는 정렬이 아니라 "여기서 칸이 나뉜다"만
    #    항상 명확히 표시하고, 실제 메일 본문에 이 문자가 등장할 일이 거의 없어 원본과도 안 헷갈린다.
    # 💡 [2026-08-28 버그 수정 ④] 위 우선순위 수정으로 CELL_BOUNDARY가 드디어 실행되긴 했지만, 실제
    #    Outlook 표는 칸 안의 글자를 <td><p>내용</p></td>처럼 Word 문단(<p>) 태그로 감싸고, 그마저도
    #    헤더 칸처럼 한 칸 안에 문단이 여러 개( <p>Achievable</p><p>Brightness Min.</p> )인 경우가
    #    흔했다. "</p>는 무조건 개행"이라는 기존 규칙을 표 안에서도 그대로 적용하면, 칸 하나의 내용이
    #    여러 줄로 쪼개지면서 "칸 경계"가 "행 경계"와 뒤섞여버려 — 정작 한 행이어야 할 칸들이 서로
    #    다른 줄에 흩어지고(예: "칸1│칸2첫줄" 한 줄, "칸2둘째줄│칸3" 다음 줄) 표를 알아볼 수 없게
    #    나왔다. 표(<table>...</table>) 안에서는 "칸 하나 = 항상 한 줄"이 되도록, 칸 내부의 문단
    #    구분(</p>, </div>, <br>)을 개행이 아니라 공백으로 합치고, 실제 줄바꿈은 오직 행 경계(</tr>)
    #    에서만 만든다 — 표 밖 본문의 <br>/<p>/<div>는 기존처럼 그대로 개행 유지.
    # 💡 [2026-08-28 버그 수정 ⑤] 표를 고치다 같이 발견 — <style> 블록(예: Outlook의 v\:*, o\:*,
    #    w\:* VML 스타일 정의)을 태그만 벗겨내고 안의 CSS 텍스트는 안 지워서, "원문 보기"/AI 분석
    #    맨 앞에 표와 무관한 CSS 잡음(v\:* {behavior:url(#default#VML);} 등)이 그대로 섞여 나왔다.
    #    JS 쪽(_mfExtractText)은 script/style을 통째로 remove()하고 있었는데 이 백엔드만 빠져 있었다.
    body = re.sub(r"(?is)<style\b.*?</style\s*>", "", body)
    body = re.sub(r"(?is)<script\b.*?</script\s*>", "", body)
    CELL_BOUNDARY = " │ "

    def _flatten_table_block(m):
        tbl = m.group(0)
        tbl = re.sub(r"(?i)<br\s*/?>", " ", tbl)
        tbl = re.sub(r"(?i)</p\s*>|</div\s*>", " ", tbl)
        tbl = re.sub(r"(?i)</t[dh]\s*>", CELL_BOUNDARY, tbl)      # 표 칸 끝 → 경계문자(다음 칸과 구분)
        tbl = re.sub(r"(?i)</tr\s*>", "\n", tbl)                  # 표 행 끝 → 줄바꿈(칸 경계와 절대 안 섞이게)
        tbl = re.sub(r"(?i)<table[^>]*>|</table\s*>", "\n", tbl)  # 표 시작/끝도 앞뒤 글과 분리
        return tbl

    body = re.sub(r">\s+<", "><", body)
    body = re.sub(r"(?is)<table\b.*?</table\s*>", _flatten_table_block, body)  # 표는 먼저 통째로 처리(중첩 표는 미지원)
    body = re.sub(r"(?i)<br\s*/?>", "\n", body)                   # 표 밖 나머지는 기존 그대로 개행
    body = re.sub(r"(?i)</p\s*>|</div\s*>", "\n", body)
    body = re.sub(r"<[^>]+>", "", body)
    # 💡 [2026-08-28 버그 수정 ⑥] 표를 고치다 같이 발견 — &nbsp;/&lt;/&gt;/&quot;/&amp; 같은 HTML
    #    엔티티를 한 번도 실제 문자로 디코딩하지 않아서 "&lt;vbernard@lnw.com&gt;"처럼 원문이 그대로
    #    노출되고 있었다. 태그를 다 벗겨낸 "뒤"에 디코딩해야 안전하다 — 미리 디코딩하면 본문에 있던
    #    "&lt;100"(=<100이라는 뜻의 문장)이 실제 "<" 문자가 되어 버려서, 그 뒤 어딘가의 ">"와 짝지어져
    #    위 태그 제거 정규식이 그 사이를 "태그"로 오인해 지워버릴 위험이 있다.
    body = html.unescape(body)
    body = re.sub(r"[ \t]{8,}", "    ", body)     # 지나치게 긴 공백(레이아웃 찌꺼기)만 축소 — 표 정렬용 짧은 공백은 보존
    # 💡 위 문단→공백 치환(</p> → " ")이 칸 내용 끝의 원래 공백과 겹치면 CELL_BOUNDARY 자체의 앞뒤
    #    공백과 합쳐져 "내용  │  다음칸"처럼 보기 싫은 이중 공백이 생긴다 — 경계문자 바로 옆의 공백만
    #    한 칸으로 정리(표 정렬용 공백은 │가 없는 다른 위치라 이 치환의 영향을 받지 않음).
    body = re.sub(r" {2,}(?=│)|(?<=│) {2,}", " ", body)
    # 💡 국내 메일 특유의 "문장. \n \n다음문장." 패턴 — 완전히 빈 줄까지 전부 제거해 간격을 촘촘하게
    #    (앞쪽 들여쓰기/표 정렬 공백은 유지하기 위해 rstrip만 하고, 빈 줄 판정에만 strip을 씀)
    #    각 행의 "마지막 칸" 뒤에는 다음 칸이 없어 경계문자( │ )가 덜렁 남으므로 같이 정리한다.
    lines = [ln.rstrip(" \t" + CELL_BOUNDARY) for ln in body.split("\n")]
    lines = [ln for ln in lines if ln.strip()]
    body = "\n".join(lines)
    # 💡 [2026-08-24] "원문 보기"용 저장 한도. AI 분석 입력은 프론트(msCallGemini)에서 이 값과 무관하게
    #    항상 별도로 2000자로 다시 잘라 쓰므로, 여기를 늘려도 AI 분석에는 영향 없음 — 저장/표시용 한도만 확장.
    return body[:15000]

def extract_attachments(msg):
    """MIME 메시지에서 첨부파일 메타데이터({name, size, type})만 추출 — 파일 내용은 읽지 않음."""
    attachments = []
    if not msg.is_multipart():
        return attachments
    for part in msg.walk():
        cd = str(part.get("Content-Disposition", ""))
        if "attachment" not in cd:
            continue
        filename = part.get_filename()
        if not filename:
            continue
        try:
            filename = decode_str(filename)
        except Exception:
            pass
        try:
            payload = part.get_payload()
            # base64 인코딩된 경우 문자열 길이로 원본 바이트 크기를 추정
            raw_len = len(payload) if isinstance(payload, (str, bytes)) else 0
            size = int(raw_len * 3 / 4) if isinstance(payload, str) else raw_len
        except Exception:
            size = 0
        attachments.append({
            "name": filename,
            "size": size,
            "type": part.get_content_type() or "application/octet-stream"
        })
    return attachments

def matches_keyword(subject, body, keyword, keyword_from='', sender='', keyword_body=''):
    def check(text, kw):
        if not kw or not kw.strip():
            return True
        keys = [k.strip().lower() for k in kw.split(",") if k.strip()]
        return any(k in text.lower() for k in keys)

    # 각 필드 독립 AND 조건 — 입력한 필드만 필터링
    if not check(subject, keyword):       return False
    if not check(sender,  keyword_from):  return False
    if not check(body,    keyword_body):  return False
    return True


# ══════════════════════════════════════════════════════════════
# 💬 Telegram 유틸
# ══════════════════════════════════════════════════════════════
def send_telegram_msg(message: str, chat_id: str = None) -> dict:
    target = chat_id or TELEGRAM_CHAT_ID
    if not TELEGRAM_TOKEN or not target:
        return {"ok": False, "error": "Token 또는 Chat ID 미설정"}
    url  = f"https://api.telegram.org/bot{TELEGRAM_TOKEN}/sendMessage"
    resp = requests.post(url, json={
        "chat_id":    target,
        "text":       message,
        "parse_mode": "HTML"
    }, timeout=10)
    return resp.json()

def get_target_chat_ids(to=None, role=None, project=None) -> list:
    if not TELEGRAM_MEMBERS:
        return [TELEGRAM_CHAT_ID] if TELEGRAM_CHAT_ID else []
    if to == "ALL":
        return [m["chat_id"] for m in TELEGRAM_MEMBERS]
    if to:
        return [m["chat_id"] for m in TELEGRAM_MEMBERS if m.get("name") == to]
    if role:
        return [m["chat_id"] for m in TELEGRAM_MEMBERS if role in m.get("roles", [])]
    if project:
        return [m["chat_id"] for m in TELEGRAM_MEMBERS
                if "ALL" in m.get("projects", []) or project in m.get("projects", [])]
    return [TELEGRAM_CHAT_ID] if TELEGRAM_CHAT_ID else []

def send_telegram_multi(message: str, chat_ids: list) -> list:
    url, results = f"https://api.telegram.org/bot{TELEGRAM_TOKEN}/sendMessage", []
    for cid in chat_ids:
        try:
            resp = requests.post(url, json={
                "chat_id": cid, "text": message, "parse_mode": "HTML"
            }, timeout=10)
            results.append({"chat_id": cid, "result": resp.json()})
        except Exception as e:
            results.append({"chat_id": cid, "error": str(e)})
    return results


# ══════════════════════════════════════════════════════════════
# 🔗 API 엔드포인트
# ══════════════════════════════════════════════════════════════

# ── OPTIONS 프리플라이트 공통 처리 ────────────────────────────
@app.before_request
def handle_options():
    if request.method == "OPTIONS":
        from flask import make_response
        resp = make_response("", 204)
        return resp

# ── 헬스체크 ─────────────────────────────────────────────────
@app.route('/health', methods=['GET'])
def health():
    return jsonify({
        'ok':      True,
        'service': 'kortek-backend-v1',
        'port':    5000,
        'smtp':    SMTP_USER  or '(미설정)',
        'pop3':    POP3_HOST,
        'telegram': '연결됨' if TELEGRAM_TOKEN else '(미설정)',
        'members': len(TELEGRAM_MEMBERS),
        'scheduleRules': len(SCHEDULE_RULES),
        'time':    datetime.now(KST).strftime('%Y-%m-%d %H:%M:%S')
    })


# ── SMTP 설정 저장 ────────────────────────────────────────────
@app.route('/config', methods=['POST'])
def update_smtp_config():
    global SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASS
    data      = request.json or {}
    SMTP_HOST = data.get('host', SMTP_HOST)
    SMTP_PORT = int(data.get('port', SMTP_PORT))
    SMTP_USER = data.get('user', SMTP_USER)
    SMTP_PASS = data.get('pass', SMTP_PASS)
    try:
        with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
            json.dump({'host': SMTP_HOST, 'port': SMTP_PORT,
                       'user': SMTP_USER, 'pass': SMTP_PASS},
                      f, ensure_ascii=False)
        return jsonify({'ok': True})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500


# ── 메일 발송 코어 (SMTP) — /send-mail 엔드포인트와 예약 발송 스케줄러가 공용으로 씀 ──
def _send_mail_core(to, subject, body, cc=''):
    """반환: (ok: bool, error: str|None, http_status: int)"""
    if not to:
        return False, '수신자 없음', 400
    if not SMTP_USER or not SMTP_PASS:
        return False, 'SMTP 미설정 — 알람 설정에서 SMTP를 입력하세요', 400
    try:
        msg            = MIMEMultipart('alternative')
        msg['From']    = SMTP_USER
        msg['To']      = to
        if cc:
            msg['Cc'] = cc
        msg['Subject'] = subject
        msg.attach(MIMEText(body, 'html', 'utf-8'))

        recipients = [r.strip() for r in to.split(',') if r.strip()]
        if cc:
            recipients += [r.strip() for r in cc.split(',') if r.strip()]

        with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=10) as smtp:
            smtp.ehlo()
            if SMTP_TLS:
                smtp.starttls()
                smtp.ehlo()
            smtp.login(SMTP_USER, SMTP_PASS)
            smtp.sendmail(SMTP_USER, recipients, msg.as_string())
        return True, None, 200
    except smtplib.SMTPAuthenticationError:
        return False, '로그인 실패: 비밀번호를 확인하세요', 401
    except smtplib.SMTPConnectError:
        return False, f'서버 연결 실패: {SMTP_HOST}:{SMTP_PORT}', 503
    except Exception as e:
        return False, str(e), 500


# ── 메일 발송 (SMTP) ──────────────────────────────────────────
@app.route('/send-mail', methods=['POST'])
def send_mail():
    data    = request.json or {}
    to      = data.get('to', '').strip()
    cc      = data.get('cc', '').strip()
    subject = data.get('subject', '').strip()
    body    = data.get('body', '').strip()

    ok, err, status = _send_mail_core(to, subject, body, cc)
    if ok:
        print(f"[OK] 메일 발송 → {to} / {subject}")
        return jsonify({'ok': True})
    return jsonify({'ok': False, 'error': err}), status


# ── 메일 수신 (POP3) ──────────────────────────────────────────
@app.route('/fetch-mail', methods=['POST'])
def fetch_mail():
    req        = request.get_json(force=True, silent=True) or {}
    mail_user  = req.get('mailUser', '')
    mail_pw    = req.get('mailPw', '')
    start_date = req.get('startDate', '')
    end_date   = req.get('endDate', '')
    keyword         = req.get('keyword', '')
    keyword_from    = req.get('keywordFrom', '')
    keyword_body    = req.get('keywordBody', '')
    max_count  = req.get('maxCount', 200)

    try:
        start_dt = datetime.strptime(start_date, "%Y-%m-%d").replace(tzinfo=KST)
        end_dt   = datetime.strptime(end_date,   "%Y-%m-%d").replace(
                       hour=23, minute=59, second=59, tzinfo=KST)

        M            = poplib.POP3(POP3_HOST, POP3_PORT, timeout=15)
        M.user(mail_user)
        M.pass_(mail_pw)
        num_messages = len(M.list()[1])
        results      = []
        scan_limit   = min(num_messages, max_count or 200)

        for i in range(num_messages, num_messages - scan_limit, -1):
            try:
                raw = b"\n".join(M.retr(i)[1])
                msg = email.message_from_bytes(raw)
                date_str = msg.get("Date", "")
                try:
                    msg_dt = parsedate_to_datetime(date_str)
                    if msg_dt.tzinfo is None:
                        msg_dt = msg_dt.replace(tzinfo=KST)
                    msg_dt = msg_dt.astimezone(KST)
                except Exception:
                    continue
                if msg_dt < start_dt:
                    break
                if msg_dt > end_dt:
                    continue
                subject = decode_str(msg.get("Subject", ""))
                sender  = decode_str(msg.get("From", ""))
                body    = extract_body(msg)
                if not matches_keyword(subject, body, keyword, keyword_from, sender, keyword_body):
                    continue

                # 💡 [우선순위 점수] 발신자가 명시한 중요도 헤더 — Outlook 등에서 "높음"으로 보낸 메일에 붙음
                importance_raw = (msg.get("Importance", "") or msg.get("X-Priority", "")).strip().lower()
                importance_high = importance_raw in ("high", "1", "1 (highest)", "2 (high)")

                # 💡 [우선순위 점수] 내가 To(직접수신)인지 Cc(참조)인지 — mail_user 계정 기준으로 판별
                # 💡 [2026-09-06 신규] 위 판별용으로만 쓰던 디코딩 결과(소문자)와 별개로, 원본 대소문자를
                #    보존한 값을 "to"/"cc"로도 함께 내려준다 — 프런트에서 AI 프롬프트의 수신자/참조 배경
                #    정보로 사용(본문에 "수신:" 줄이 없는 메일의 폴백 근거). 헤더값은 보통 이메일 주소라
                #    "받는사람" 표시로는 body에 적힌 사람 이름보다 덜 직관적이지만, 문맥 단서가 전혀 없을
                #    때 "수신자 미지정"보다는 실제 주소를 보여주는 편이 낫다는 판단.
                to_raw    = decode_str(msg.get("To", ""))
                cc_raw    = decode_str(msg.get("Cc", ""))
                to_header = to_raw.lower()
                cc_header = cc_raw.lower()
                my_addr   = mail_user.lower()
                is_to_me  = my_addr in to_header
                is_cc_me  = (not is_to_me) and (my_addr in cc_header)

                results.append({
                    "subject":     subject,
                    "sender":      sender,
                    "to":          to_raw,
                    "cc":          cc_raw,
                    "date":        msg_dt.strftime("%Y-%m-%d %H:%M"),
                    "body":        body,
                    "fileName":    f"{msg_dt.strftime('%Y%m%d')}_{subject[:20]}.eml",
                    "importance":  importance_high,           # true/false
                    "isToMe":      is_to_me,                  # true/false
                    "isCcMe":      is_cc_me,                  # true/false
                    "attachments": extract_attachments(msg)   # [{name,size,type}, ...]
                })
            except Exception:
                continue

        M.quit()
        return jsonify({
            "status":  "success",
            "count":   len(results),
            "keyword": keyword or "(전체)",
            "userId":  mail_user.split("@")[0] if mail_user else "",
            "data":    results
        })
    except poplib.error_proto as e:
        return jsonify({"status": "error", "message": f"POP3 인증 실패: {str(e)}"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)})


# ── Telegram 메시지 발송 ──────────────────────────────────────
@app.route('/send-telegram', methods=['POST'])
def send_telegram_api():
    req     = request.get_json(force=True, silent=True) or {}
    message = req.get('message', '')
    to      = req.get('to', '')
    role    = req.get('role', '')
    project = req.get('project', '')
    chat_id = req.get('chatId', '')

    if not message:
        return jsonify({'status': 'error', 'message': 'message 필드 필요'})
    if not TELEGRAM_TOKEN:
        return jsonify({'status': 'error', 'message': 'Telegram Token 미설정 — 알람 설정 > Telegram 탭에서 입력하세요'})

    try:
        if chat_id:
            chat_ids = [chat_id]
        else:
            chat_ids = get_target_chat_ids(
                to=to or None, role=role or None, project=project or None)
        if not chat_ids:
            return jsonify({'status': 'error', 'message': '발송 대상 없음'})

        results = send_telegram_multi(message, chat_ids)
        ok_cnt  = sum(1 for r in results if r.get('result', {}).get('ok'))
        return jsonify({'status': 'success', 'sent': ok_cnt, 'total': len(results), 'results': results})
    except Exception as e:
        return jsonify({'status': 'error', 'message': str(e)})


# ── Telegram 설정 저장/조회 ───────────────────────────────────
@app.route('/telegram/config', methods=['GET', 'POST'])
def telegram_config_api():
    global TELEGRAM_TOKEN, TELEGRAM_CHAT_ID

    if request.method == 'GET':
        token = TELEGRAM_TOKEN
        return jsonify({
            'status':          'success',
            'token_set':       bool(token),
            'token_preview':   (token[:10] + '...') if token else '(미설정)',
            'default_chat_id': TELEGRAM_CHAT_ID,
            'members_count':   len(TELEGRAM_MEMBERS)
        })

    data            = request.json or {}
    TELEGRAM_TOKEN   = data.get('token',           TELEGRAM_TOKEN)
    TELEGRAM_CHAT_ID = data.get('default_chat_id', TELEGRAM_CHAT_ID)
    try:
        cfg = load_tg_config()
        cfg['token']           = TELEGRAM_TOKEN
        cfg['default_chat_id'] = TELEGRAM_CHAT_ID
        save_tg_config(cfg)
        return jsonify({'ok': True, 'message': 'Telegram 설정 저장 완료'})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500


# ── 팀원 관리 ─────────────────────────────────────────────────
@app.route('/telegram/members', methods=['GET', 'POST', 'DELETE'])
def telegram_members_api():
    global TELEGRAM_MEMBERS

    if request.method == 'GET':
        return jsonify({'status': 'success', 'count': len(TELEGRAM_MEMBERS),
                        'members': TELEGRAM_MEMBERS})

    if request.method == 'POST':
        data     = request.json or {}
        name     = data.get('name',     '').strip()
        chat_id  = data.get('chat_id',  '').strip()
        email_   = data.get('email',    '').strip()
        roles    = data.get('roles',    [])
        projects = data.get('projects', [])
        if not name or not chat_id:
            return jsonify({'ok': False, 'error': 'name, chat_id 필수'}), 400

        updated = False
        for m in TELEGRAM_MEMBERS:
            if m['name'] == name:
                m.update({'chat_id': chat_id, 'email': email_,
                          'roles': roles, 'projects': projects})
                updated = True
                break
        if not updated:
            TELEGRAM_MEMBERS.append({'name': name, 'chat_id': chat_id,
                                     'email': email_, 'roles': roles, 'projects': projects})
        cfg = load_tg_config()
        cfg['members'] = TELEGRAM_MEMBERS
        save_tg_config(cfg)
        return jsonify({'ok': True, 'message': f"{'업데이트' if updated else '등록'} 완료: {name}"})

    if request.method == 'DELETE':
        data   = request.json or {}
        name   = data.get('name', '').strip()
        before = len(TELEGRAM_MEMBERS)
        TELEGRAM_MEMBERS = [m for m in TELEGRAM_MEMBERS if m['name'] != name]
        cfg = load_tg_config()
        cfg['members'] = TELEGRAM_MEMBERS
        save_tg_config(cfg)
        removed = before - len(TELEGRAM_MEMBERS)
        return jsonify({'ok': True, 'message': f'{name} 삭제 완료' if removed else '해당 이름 없음'})


# ── Telegram 발송 테스트 ──────────────────────────────────────
@app.route('/telegram/test', methods=['POST'])
def telegram_test():
    if not TELEGRAM_TOKEN:
        return jsonify({'ok': False, 'error': 'Token 미설정'})
    data    = request.json or {}
    chat_id = data.get('chat_id', TELEGRAM_CHAT_ID)
    try:
        result = send_telegram_msg(
            f"✅ <b>KORTEK Gantt PM</b>\n"
            f"Telegram 연결 테스트 성공!\n"
            f"서버 시간: {datetime.now(KST).strftime('%Y-%m-%d %H:%M:%S')}",
            chat_id
        )
        if result.get('ok'):
            return jsonify({'ok': True,  'message': '테스트 메시지 발송 완료'})
        else:
            return jsonify({'ok': False, 'error': str(result)})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)})


# ══════════════════════════════════════════════════════════════
# ⏰ 예약 발송(반복 규칙) — 규칙 등록/조회/삭제 + 백그라운드 스케줄러
# ══════════════════════════════════════════════════════════════
#
# 규칙(rule) 스키마:
#   id, type('notice'|'alarm'), title, message,
#   recipients: [{name, email, telegramId, emailOn, tgOn}, ...],
#   dateMode('range'|'specific'),
#     range  → startDate, endDate, dayInterval(N일마다)
#     specific → specificDates: ["YYYY-MM-DD", ...]
#   hourStart, hourEnd ("HH:MM", 기본 09:00~21:00), hourInterval(N시간마다),
#   enabled, createdAt, lastFiredBucket(중복발송 방지용, 서버가 내부적으로 기록)
#
# [2026-08-31 신규] type='alarm'(업무별 개별 알람) 전용 추가 필드 — 공지(notice)와 달리
# "저장 시점 스냅샷"이 아니라 "발송 직전 구글드라이브에서 해당 업무의 최신 상태를 다시 읽어서"
# 담당자/완료예정일/업무내용이 바뀌었으면 바뀐 대로 발송한다(_fire_alarm_rule 참고).
#   driveFileId          — 이 업무가 속한 프로젝트의 구글드라이브 파일 ID (window.currentDriveFileId)
#   rowIdx               — globalData 배열 안에서 이 업무 행의 인덱스
#   taskNameSnapshot      — 목록 표시용(실제 발송 내용에는 안 씀, 매번 최신 이름으로 다시 계산됨)
#   ccRecipientsSnapshot  — 수신 대상(기본수신/개별수신) 명단은 브라우저에만 있어 서버가 못 보므로,
#                           규칙 저장 시점의 명단을 [{name, email, telegramId, emailOn, tgOn}, ...]
#                           그대로 스냅샷 — recipients와 달리 자주 안 바뀔 값이라 실용적 절충.
#                           [2026-08-31] 이메일만 담던 ccMailsSnapshot(문자열)을 대체 — 텔레그램도 발송.
#   allowedExternalDomainsSnapshot — 위와 동일한 이유로 스냅샷하는 "외부 도메인 발송 허용" 목록

@app.route('/schedule', methods=['GET', 'POST'])
def schedule_api():
    global SCHEDULE_RULES
    if request.method == 'GET':
        with _SCHEDULE_LOCK:
            return jsonify({'ok': True, 'rules': SCHEDULE_RULES})

    data     = request.json or {}
    rule_id  = data.get('id') or str(uuid.uuid4())
    rule_type = data.get('type', 'notice')
    rule = {
        'id':            rule_id,
        'type':          rule_type,
        'title':         (data.get('title') or '').strip(),
        'message':       data.get('message', ''),
        'recipients':    data.get('recipients', []),
        'dateMode':      data.get('dateMode', 'range'),
        'startDate':     data.get('startDate', ''),
        'endDate':       data.get('endDate', ''),
        'specificDates': data.get('specificDates', []),
        'dayInterval':   max(1, int(data.get('dayInterval', 1) or 1)),
        'hourStart':     data.get('hourStart', '09:00'),
        'hourEnd':       data.get('hourEnd', '21:00'),
        'hourInterval':  max(0.25, float(data.get('hourInterval', 1) or 1)),
        'enabled':       bool(data.get('enabled', True)),
        'createdAt':     data.get('createdAt') or datetime.now(KST).isoformat(),
    }
    if rule_type == 'alarm':
        # 💡 업무별 알람은 "지금" 제목/내용을 저장하는 게 아니라, driveFileId+rowIdx로 어느 업무인지만
        #    기억해뒀다가 발송 직전에 매번 구글드라이브에서 최신 상태를 다시 읽는다(_fire_alarm_rule).
        rule['driveFileId']                     = data.get('driveFileId', '')
        rule['rowIdx']                          = data.get('rowIdx')
        rule['taskNameSnapshot']                = data.get('taskNameSnapshot', '')
        rule['ccRecipientsSnapshot']            = data.get('ccRecipientsSnapshot', []) or []
        rule['allowedExternalDomainsSnapshot']  = data.get('allowedExternalDomainsSnapshot', []) or []
        if not rule['title']:
            rule['title'] = rule['taskNameSnapshot'] or '(제목 없음)'
        if not rule['driveFileId'] or rule['rowIdx'] is None:
            return jsonify({'ok': False, 'error': 'driveFileId/rowIdx가 필요합니다 (알람 규칙)'}), 400
    if not rule['title']:
        return jsonify({'ok': False, 'error': '제목이 필요합니다'}), 400

    with _SCHEDULE_LOCK:
        existing = next((r for r in SCHEDULE_RULES if r.get('id') == rule_id), None)
        if existing:
            # 💡 이력(마지막 발송 버킷)은 유지한 채 나머지 설정만 덮어씀 — 안 그러면 저장할 때마다
            #    "이미 오늘 보냈는지" 기록이 사라져 같은 시각에 중복 발송될 수 있음.
            rule['lastFiredBucket'] = existing.get('lastFiredBucket')
            rule['lastFiredAt']     = existing.get('lastFiredAt')
            SCHEDULE_RULES[SCHEDULE_RULES.index(existing)] = rule
        else:
            SCHEDULE_RULES.append(rule)
        save_schedule_rules(SCHEDULE_RULES)
    return jsonify({'ok': True, 'id': rule_id})


@app.route('/schedule/<rule_id>', methods=['DELETE'])
def schedule_delete_api(rule_id):
    global SCHEDULE_RULES
    with _SCHEDULE_LOCK:
        before = len(SCHEDULE_RULES)
        SCHEDULE_RULES = [r for r in SCHEDULE_RULES if r.get('id') != rule_id]
        removed = before - len(SCHEDULE_RULES)
        save_schedule_rules(SCHEDULE_RULES)
    return jsonify({'ok': True, 'removed': removed > 0})


# ══════════════════════════════════════════════════════════════
# 📂 업무별(개별 태스크) 알람용 — 구글드라이브 실시간 조회 (서비스계정, 읽기전용)
#    GANTT_CHART_V02_Color.html은 구글시트가 아니라, 이 폴더(GANTT_SHARED_FOLDER_ID) 안의
#    JSON 파일 1개(프로젝트당 1개)에 전체 데이터를 저장한다 — js/04a~04k-core-app-*.js 참고.
#    여기서는 그 JSON을 그대로 읽어 collectAlarmItems()(js/22c-summary-mctable-core2.js:3)와
#    동등한 로직으로 딱 한 업무(행)의 "지금 이 순간" 상태를 재계산한다.
# ══════════════════════════════════════════════════════════════
_drive_creds = None


def _drive_get_access_token() -> str:
    """서비스계정 JSON 키로 구글드라이브 접근 토큰(읽기전용)을 발급/갱신해서 돌려준다."""
    global _drive_creds
    if not GOOGLE_AUTH_OK:
        raise RuntimeError('google-auth 미설치 — pip install google-auth')
    if not os.path.exists(GOOGLE_SERVICE_ACCOUNT_FILE):
        raise RuntimeError(f'서비스계정 키 파일을 찾을 수 없습니다: {GOOGLE_SERVICE_ACCOUNT_FILE}')
    if _drive_creds is None:
        _drive_creds = _google_service_account.Credentials.from_service_account_file(
            GOOGLE_SERVICE_ACCOUNT_FILE,
            scopes=['https://www.googleapis.com/auth/drive.readonly']
        )
    if not _drive_creds.valid:
        _drive_creds.refresh(_GoogleAuthRequest())
    return _drive_creds.token


def _drive_find_child(parent_id: str, name: str, mime_type: str = None):
    """parent_id 폴더 바로 밑에서 이름이 name인 파일/폴더 1개의 id를 찾는다 (없으면 None)."""
    token = _drive_get_access_token()
    q = f"name='{name}' and trashed=false and '{parent_id}' in parents"
    if mime_type:
        q += f" and mimeType='{mime_type}'"
    resp = requests.get(
        'https://www.googleapis.com/drive/v3/files',
        params={'q': q, 'supportsAllDrives': 'true', 'includeItemsFromAllDrives': 'true', 'fields': 'files(id,name)'},
        headers={'Authorization': f'Bearer {token}'}, timeout=15
    )
    resp.raise_for_status()
    files = (resp.json() or {}).get('files') or []
    return files[0]['id'] if files else None


def _drive_read_json_file(file_id: str) -> dict:
    """파일 id로 JSON 파일 내용을 그대로 읽어온다 (alt=media)."""
    token = _drive_get_access_token()
    resp = requests.get(
        f'https://www.googleapis.com/drive/v3/files/{file_id}',
        params={'alt': 'media', 'supportsAllDrives': 'true'},
        headers={'Authorization': f'Bearer {token}'}, timeout=20
    )
    resp.raise_for_status()
    return resp.json()


_app_config_folder_cache = {'id': None}


def _drive_get_app_config_folder_id():
    """'App_Config' 하위 폴더 id — AddressBook_Shared.json 등이 들어있다 (js: getOrCreateConfigFolder)."""
    if _app_config_folder_cache['id']:
        return _app_config_folder_cache['id']
    fid = _drive_find_child(GANTT_SHARED_FOLDER_ID, 'App_Config', mime_type='application/vnd.google-apps.folder')
    _app_config_folder_cache['id'] = fid
    return fid


_address_book_cache = {'data': [], 'loadedAt': 0.0}


def _drive_load_address_book(force: bool = False) -> list:
    """AddressBook_Shared.json(이름→이메일/텔레그램ID, 모든 프로젝트 공유) — 자주 안 바뀌므로 캐시."""
    now = time.time()
    if not force and (now - _address_book_cache['loadedAt']) < ADDRESS_BOOK_CACHE_TTL_SEC:
        return _address_book_cache['data']
    try:
        folder_id = _drive_get_app_config_folder_id() or GANTT_SHARED_FOLDER_ID
        file_id = _drive_find_child(folder_id, 'AddressBook_Shared.json')
        if not file_id and folder_id != GANTT_SHARED_FOLDER_ID:
            file_id = _drive_find_child(GANTT_SHARED_FOLDER_ID, 'AddressBook_Shared.json')  # 구버전 위치 폴백
        if file_id:
            data = _drive_read_json_file(file_id)
            _address_book_cache['data'] = data.get('addressBook') or []
            _address_book_cache['loadedAt'] = now
    except Exception as e:
        print(f"[업무알람] 주소록 로드 실패(이전 캐시로 계속 진행): {e}")
    return _address_book_cache['data']


# ── 이름 매칭 유틸 (js/22b-summary-mctable-core1.js의 _addrSplitNames/_addrStripTitleSuffix/
#    _addrFindByName과 동일 로직 — 메일 본문에서 뽑힌 발신/수신인 이름을 주소록과 매칭하기 위함) ──
ADDR_KO_TITLE_WORDS = [
    '회장', '부회장', '사장', '부사장', '대표', '전무', '상무', '이사', '감사',
    '본부장', '소장', '센터장', '실장', '팀장', '파트장', '그룹장', '랩장',
    '수석', '책임', '선임', '주임', '매니저', '대리', '과장', '차장', '부장', '사원', '연구원'
]
ADDR_EN_TITLE_WORDS = ['mr', 'mrs', 'ms', 'miss', 'dr', 'prof', 'manager', 'director', 'leader', 'president', 'vp', 'ceo', 'cto', 'coo']
KORTEK_INTERNAL_DOMAIN = 'kortek.co.kr'


def _addr_strip_title_suffix(name: str) -> str:
    n = (name or '').strip()
    changed = True
    while changed:
        changed = False
        before = n
        n = re.sub(r'\s*(님|씨)\s*$', '', n).strip()
        for t in ADDR_KO_TITLE_WORDS:
            if n.endswith(t) and len(n) > len(t):
                n = n[:-len(t)].strip(); break
            if n.startswith(t) and len(n) > len(t):
                n = n[len(t):].strip(); break
        for t in ADDR_EN_TITLE_WORDS:
            re_suf = re.compile(r'[.,]?\s*' + t + r'\.?$', re.IGNORECASE)
            re_pre = re.compile(r'^' + t + r'\.?\s*', re.IGNORECASE)
            if re_suf.search(n):
                n = re_suf.sub('', n).strip(); break
            if re_pre.search(n):
                n = re_pre.sub('', n).strip(); break
        if n != before:
            changed = True
    return n


def _addr_split_names(s: str) -> list:
    """"정민희/임희철", "박용훈 외 다수" 같은 패턴도 개별 이름으로 분리"""
    if not s:
        return []
    out = []
    for p in re.split(r'[,，/]', str(s)):
        p = re.sub(r'\s*외\s*(\d+\s*(명|인)?|다수)?\s*$', '', p.strip()).strip()
        if p:
            out.append(p)
    return out


def _addr_find_by_name(address_book: list, name: str):
    """한글 이름 정확일치 → 영문 이름 정확일치 → (둘 다 실패 시) 직함/존칭 뗀 이름으로 재시도"""
    if not name:
        return None

    def try_exact(n):
        if not n:
            return None
        for p in address_book:
            if (p.get('name') or '').strip() == n:
                return p
        for p in address_book:
            if (p.get('nameEn') or '').strip().lower() == n.lower():
                return p
        return None

    trimmed = str(name).strip()
    found = try_exact(trimmed)
    if found:
        return found
    stripped = _addr_strip_title_suffix(trimmed)
    if stripped and stripped != trimmed:
        found = try_exact(stripped)
    return found


def _lookup_email(address_book: list, name: str) -> str:
    if not name:
        return ''
    trimmed = name.strip()
    if re.match(r'^[^\s@]+@[^\s@]+\.[^\s@]+$', trimmed):  # 이름 자리에 이메일이 이미 들어있는 경우
        return trimmed
    found = _addr_find_by_name(address_book, trimmed)
    return (found.get('email') or '') if found else ''


def _is_alarm_domain_allowed(email: str, allowed_external_domains: list) -> bool:
    """@kortek.co.kr은 항상 허용, 그 외 도메인은 규칙 저장 시점에 스냅샷된 허용 목록에 있어야 허용"""
    if not email:
        return False
    m = re.search(r'@([^@\s]+)$', email.strip().lower())
    if not m:
        return False
    if m.group(1) == KORTEK_INTERNAL_DOMAIN:
        return True
    return m.group(1) in (allowed_external_domains or [])


def _build_alarm_task_snapshot(project_data: dict, row_idx: int):
    """딱 한 업무(row_idx)의 "지금 이 순간" 상태를 계산한다 — collectAlarmItems()의 파이썬 버전.
       프로젝트 JSON은 저장 시점에 globalData의 모든 _밑줄필드(_level, _origDev 등)를 그대로
       보존하므로(js/04c-core-app-mail-pipeline.js:628 serializedGlobalData), 화면 렌더링 로직을 다시 짤 필요
       없이 저장된 값을 그대로 읽으면 된다. 대상이 아니면(알림 꺼짐/완료예정일 없음/행 없음) None."""
    global_data = project_data.get('globalData') or []
    if row_idx < 0 or row_idx >= len(global_data):
        return None
    row_obj = global_data[row_idx]
    if not isinstance(row_obj, dict):
        return None
    if not row_obj.get('_알림'):  # 알람이 등록되지 않은 업무
        return None

    cols    = row_obj.get('data') or []
    col_idx = project_data.get('colIdx') or {}

    def col(name):
        idx = col_idx.get(name)
        if idx is None or idx < 0 or idx >= len(cols):
            return ''
        return cols[idx]

    # 완료 예정일 — plan 컬럼(YYYY-MM-DD), 자동(🔓) 모드 행은 비어있고 _calcPlanTs만 있어 폴백 처리
    due_raw = str(col('plan') or '').strip()
    due_date = None
    if due_raw and due_raw != '-':
        try:
            due_date = datetime.strptime(due_raw[:10], '%Y-%m-%d').date()
        except Exception:
            due_date = None
    if due_date is None and row_obj.get('_calcPlanTs'):
        try:
            due_date = datetime.fromtimestamp(row_obj['_calcPlanTs'] / 1000, KST).date()
        except Exception:
            due_date = None
    if due_date is None:
        return None

    today     = datetime.now(KST).date()
    diff_days = (due_date - today).days
    due_str   = due_date.strftime('%Y-%m-%d')

    # 업무명 — 레벨별 원본 이름(_origDev/_origT1~4) 우선, 없으면 그 레벨의 WBS 컬럼값으로 폴백
    level = row_obj.get('_level')
    orig_by_level = {
        0: row_obj.get('_origDev'), 1: row_obj.get('_origT1'), 2: row_obj.get('_origT2'),
        3: row_obj.get('_origT3'), 4: row_obj.get('_origT4'),
    }.get(level)
    wbs_col_name = {0: 'devStage', 1: 'taskType1', 2: 'taskType2', 3: 'taskType3', 4: 'taskType4'}.get(level, 'wbs')
    task_name = str(orig_by_level or col(wbs_col_name) or '').strip() or '-'
    task_name = re.sub(r'^🌐\s*', '', task_name)

    content_raw = str(col('content') or '').strip()

    # 💡 [2026-09-01 신규] 알람 일정 모달의 "업무 정보" 편집칸에서 이 알람만을 위해 제목/내용을
    #    덮어썼으면(js/22c-summary-mctable-core2.js의 row._알림제목오버라이드/_알림내용오버라이드,
    #    saveAlarmSchedule/_asSaveRecurRule에서 저장) 발신인/수신인 추출은 원본 content_raw 그대로 두고
    #    실제 메일에 쓰이는 task_name/content만 덮어쓴다(collectAlarmItems의 파이썬 버전이라 동일하게 처리).
    title_override = row_obj.get('_알림제목오버라이드')
    if title_override and str(title_override).strip():
        task_name = str(title_override).strip()
    content_override = row_obj.get('_알림내용오버라이드')
    if content_override and str(content_override).strip():
        content_raw_for_mail = str(content_override).strip()
    else:
        content_raw_for_mail = content_raw

    # [발신인→수신인] 패턴 — 없으면 담당자(assignee) 컬럼을 발신인으로 취급
    arrow_match = re.search(r'\[([^\]→]+)→((?:\[[^\]]*\]|[^\]])+)\]', content_raw)

    def strip_tag(s):
        return re.sub(r'^\[[^\]]*\]\s*', '', (s or '').strip()).strip()

    sender_raw   = strip_tag(arrow_match.group(1)) if arrow_match else str(col('assignee') or '-').strip()
    receiver_raw = strip_tag(arrow_match.group(2)) if arrow_match else ''

    address_book   = _drive_load_address_book()
    sender_names   = _addr_split_names(sender_raw)
    receiver_names = _addr_split_names(receiver_raw)
    all_people     = list(dict.fromkeys([n for n in (sender_names + receiver_names) if n]))

    assignee        = ', '.join(sender_names) or '-'
    assignee_email  = ','.join(dict.fromkeys([e for e in (_lookup_email(address_book, n) for n in sender_names) if e]))
    receiver_str    = ', '.join(receiver_names) or '-'
    receiver_emails = list(dict.fromkeys([e for e in (_lookup_email(address_book, n) for n in receiver_names) if e]))

    status_val = str(col('status') or '').strip() if col_idx.get('status', -1) != -1 else ''

    return {
        'taskName': task_name, 'status': status_val,
        'assignee': assignee, 'assigneeEmail': assignee_email,
        'receiverStr': receiver_str, 'receiverEmails': receiver_emails,
        'allPeopleNames': all_people,
        'dueStr': due_str, 'diffDays': diff_days, 'content': content_raw_for_mail,
    }


def _rule_active_today(rule: dict, today) -> bool:
    """오늘이 이 규칙의 발송 대상 날짜인지(기간+N일마다, 또는 특정 날짜 목록)"""
    if rule.get('dateMode') == 'specific':
        return today.strftime('%Y-%m-%d') in (rule.get('specificDates') or [])
    start_s, end_s = rule.get('startDate'), rule.get('endDate')
    if not start_s or not end_s:
        return False
    try:
        start = datetime.strptime(start_s, '%Y-%m-%d').date()
        end   = datetime.strptime(end_s,   '%Y-%m-%d').date()
    except Exception:
        return False
    if today < start or today > end:
        return False
    interval = max(1, int(rule.get('dayInterval', 1) or 1))
    return (today - start).days % interval == 0


def _rule_today_buckets(rule: dict) -> list:
    """오늘 이 규칙이 울려야 할 시각들을 자정 기준 분(minute) 목록으로 반환 (시간창 + N시간마다)"""
    try:
        sh, sm = map(int, (rule.get('hourStart') or '09:00').split(':'))
        eh, em = map(int, (rule.get('hourEnd')   or '21:00').split(':'))
    except Exception:
        sh, sm, eh, em = 9, 0, 21, 0
    start_min = sh * 60 + sm
    end_min   = eh * 60 + em
    step_min  = max(15, int(round(float(rule.get('hourInterval', 1) or 1) * 60)))
    buckets, t = [], start_min
    while t <= end_min:
        buckets.append(t)
        t += step_min
    return buckets


def _fire_notice_rule(rule: dict):
    """공지(type='notice') 발송 — 규칙 저장 시점에 고정된 제목/내용을 등록된 수신자 전원에게 그대로 발송"""
    title      = rule.get('title', '')
    message    = rule.get('message', '')
    recipients = rule.get('recipients', []) or []
    subject    = f"[예약 발송] {title}"
    body_html  = '<div style="white-space:pre-wrap; font-family:\'맑은 고딕\',sans-serif;">' \
                 + html.escape(message).replace('\n', '<br>') + '</div>'
    for r in recipients:
        try:
            if r.get('emailOn') and r.get('email'):
                ok, err, _ = _send_mail_core(r['email'], subject, body_html)
                if not ok:
                    print(f"[예약발송 실패-메일] {r.get('name','')} <{r['email']}>: {err}")
        except Exception as e:
            print(f"[예약발송 예외-메일] {r.get('name','')}: {e}")
        try:
            if r.get('tgOn') and r.get('telegramId'):
                res = send_telegram_msg(f"📢 <b>{html.escape(title)}</b>\n{html.escape(message)}", r['telegramId'])
                if not res.get('ok'):
                    print(f"[예약발송 실패-텔레그램] {r.get('name','')}: {res}")
        except Exception as e:
            print(f"[예약발송 예외-텔레그램] {r.get('name','')}: {e}")
    print(f"[예약발송] '{title}' → 수신자 {len(recipients)}명 처리 ({datetime.now(KST).strftime('%Y-%m-%d %H:%M:%S')})")


def _fire_alarm_rule(rule: dict):
    """업무별 알람(type='alarm') 발송 — 저장된 내용이 아니라, 발송 직전 구글드라이브에서 해당 업무의
       최신 담당자/완료예정일/업무내용을 다시 읽어 그 시점 값으로 발송한다 (_build_alarm_task_snapshot)."""
    drive_file_id = rule.get('driveFileId')
    row_idx       = rule.get('rowIdx')
    label         = rule.get('taskNameSnapshot') or rule.get('title') or '(제목 없음)'
    if not drive_file_id or row_idx is None:
        print(f"[업무알람 실패] driveFileId/rowIdx가 없는 규칙입니다: {label}")
        return
    try:
        project_data = _drive_read_json_file(drive_file_id)
    except Exception as e:
        print(f"[업무알람 실패] '{label}' — 프로젝트 파일 조회 실패: {e}")
        return

    try:
        snap = _build_alarm_task_snapshot(project_data, int(row_idx))
    except Exception as e:
        print(f"[업무알람 실패] '{label}' — 업무 상태 계산 중 오류: {e}")
        return
    if not snap:
        print(f"[업무알람 건너뜀] '{label}' — 업무를 찾을 수 없거나(삭제됨) 알림이 꺼져있거나 완료예정일이 없습니다.")
        return

    allowed_domains = rule.get('allowedExternalDomainsSnapshot') or []
    assignee_email  = snap['assigneeEmail'] if _is_alarm_domain_allowed(snap['assigneeEmail'], allowed_domains) else ''
    receiver_emails = [e for e in snap['receiverEmails'] if _is_alarm_domain_allowed(e, allowed_domains)]
    to_email = ','.join(dict.fromkeys([e for e in ([assignee_email] + receiver_emails) if e]))
    # 💡 [2026-08-31] 수신 대상(기본수신/개별수신) 스냅샷 — 이메일 켜져있고(emailOn) 도메인 허용된
    #    사람만 이메일 cc로, 텔레그램 켜져있는(tgOn) 사람은 아래에서 별도로 텔레그램 발송
    cc_recipients = rule.get('ccRecipientsSnapshot') or []
    cc_mails = ','.join(dict.fromkeys(
        r.get('email', '') for r in cc_recipients
        if r.get('emailOn') and r.get('email') and _is_alarm_domain_allowed(r['email'], allowed_domains)
    ))

    project_meta = project_data.get('projectMeta') or {}
    proj_title   = ' > '.join([x for x in [project_meta.get('고객사'), project_meta.get('고객모델명')] if x]) or '프로젝트'
    d = snap['diffDays']
    d_day_plain = 'D-Day' if d == 0 else (f'D+{abs(d)}' if d < 0 else f'D-{d}')

    if not to_email:
        print(f"[업무알람 건너뜀] '{snap['taskName']}' — 발송 가능한 이메일이 없습니다(담당자/수신인 이메일 미등록 또는 외부 도메인 차단).")
    else:
        subject = f'[Gantt 알람] {proj_title} — "{snap["taskName"]}" 완료일 {d_day_plain}'
        content_html = html.escape(snap['content']).replace('\n', '<br>') if snap['content'] else ''
        content_row = (
            f'<tr><td style="padding:6px 12px; background:#f0f8ff; font-weight:bold; vertical-align:top; border:1px solid #dcdde1; width:90px; white-space:nowrap;">내용</td>'
            f'<td style="padding:6px 12px; white-space:normal; border:1px solid #dcdde1; word-break:break-word;">{content_html}</td></tr>'
        ) if content_html else ''
        body = f"""<div style="font-family:'맑은 고딕',sans-serif; font-size:14px; color:#333;">
  <p><b>{html.escape(snap['assignee'])}님께</b></p>
  <p>아래 업무의 완료 예정일이 <b style="color:#e74c3c;">{d_day_plain}일 ({snap['dueStr']})</b>입니다.</p>
  <table style="border-collapse:collapse; margin:12px 0; border:1px solid #dcdde1; width:100%; max-width:1000px;">
    <tr><td style="padding:6px 12px; background:#f0f8ff; font-weight:bold; border:1px solid #dcdde1; width:90px; white-space:nowrap;">프로젝트</td><td style="padding:6px 12px; border:1px solid #dcdde1; word-break:break-word;">{html.escape(proj_title)}</td></tr>
    <tr><td style="padding:6px 12px; background:#f0f8ff; font-weight:bold; border:1px solid #dcdde1; width:90px; white-space:nowrap;">업무</td><td style="padding:6px 12px; border:1px solid #dcdde1; word-break:break-word;">{html.escape(snap['taskName'])}</td></tr>
    {content_row}
    <tr><td style="padding:6px 12px; background:#f0f8ff; font-weight:bold; border:1px solid #dcdde1; width:90px; white-space:nowrap;">담당자</td><td style="padding:6px 12px; border:1px solid #dcdde1; word-break:break-word;">{html.escape(snap['assignee'])}</td></tr>
    <tr><td style="padding:6px 12px; background:#f0f8ff; font-weight:bold; border:1px solid #dcdde1; width:90px; white-space:nowrap;">완료 예정일</td><td style="padding:6px 12px; color:#e74c3c; border:1px solid #dcdde1;"><b>{snap['dueStr']}</b></td></tr>
  </table>
  <p style="color:#888; font-size:12px;">본 메일은 Gantt Chart 예약 발송(기간·반복) 알람에서 자동 발송되었습니다. 담당자/완료예정일/내용은 발송 시점 기준 최신 상태입니다.</p>
</div>"""
        ok, err, _ = _send_mail_core(to_email, subject, body, cc_mails)
        if not ok:
            print(f"[업무알람 실패-메일] '{snap['taskName']}': {err}")
        else:
            print(f"[업무알람] '{snap['taskName']}' → {to_email} 발송 완료 ({datetime.now(KST).strftime('%Y-%m-%d %H:%M:%S')})")

    # 텔레그램 — 담당자/수신인 중 (이메일 미등록이라 판단 불가하거나) 도메인이 허용된 사람만
    try:
        address_book = _drive_load_address_book()
        tg_msg = f"📌 [Gantt 알람] {proj_title}\n업무: {snap['taskName']}\n담당: {snap['assignee']}\n기한: {snap['dueStr']} ({d_day_plain})"
        if snap['content']:
            tg_msg += '\n내용: ' + snap['content'].replace('\n', ' ')[:2000]
        sent_chat_ids = set()

        def _tg_send_once(chat_id, label):
            if not chat_id or chat_id in sent_chat_ids:
                return
            sent_chat_ids.add(chat_id)
            res = send_telegram_msg(tg_msg, chat_id)
            if not res.get('ok'):
                print(f"[업무알람 실패-텔레그램] {label}: {res}")

        for name in snap['allPeopleNames']:
            person = _addr_find_by_name(address_book, name)
            if not person or not person.get('telegramId'):
                continue
            email = person.get('email') or ''
            if email and not _is_alarm_domain_allowed(email, allowed_domains):
                continue
            _tg_send_once(person['telegramId'], name)

        # 💡 수신 대상(기본수신/개별수신)에 텔레그램이 켜진 사람도 함께 발송
        for r in cc_recipients:
            if not r.get('tgOn') or not r.get('telegramId'):
                continue
            email = r.get('email') or ''
            if email and not _is_alarm_domain_allowed(email, allowed_domains):
                continue
            _tg_send_once(r['telegramId'], r.get('name', ''))
    except Exception as e:
        print(f"[업무알람 텔레그램 예외] {snap['taskName']}: {e}")


def _fire_rule(rule: dict):
    """규칙 종류에 따라 실제 발송 로직을 분기한다."""
    if rule.get('type') == 'alarm':
        _fire_alarm_rule(rule)
    else:
        _fire_notice_rule(rule)


def _scheduler_tick():
    """1분마다 호출 — 오늘 발송해야 할 규칙을 찾아 발송하고 lastFiredBucket을 기록한다.
       서버가 꺼져 있던 동안 지나간 시각은 유예시간(GRACE_MINUTES)을 넘기면 그냥 건너뛴다
       (밀린 발송을 몰아서 보내지 않음 — 사용자 요청사항)."""
    GRACE_MINUTES = 2  # 체크 주기(1분)보다 약간 넉넉하게
    now         = datetime.now(KST)
    today       = now.date()
    today_str   = today.strftime('%Y-%m-%d')
    cur_minutes = now.hour * 60 + now.minute

    with _SCHEDULE_LOCK:
        rules_snapshot = list(SCHEDULE_RULES)

    fired_any = False
    for rule in rules_snapshot:
        if not rule.get('enabled', True):
            continue
        try:
            if not _rule_active_today(rule, today):
                continue
            for bm in _rule_today_buckets(rule):
                if bm > cur_minutes or (cur_minutes - bm) > GRACE_MINUTES:
                    continue  # 아직 안 됐거나, 유예시간 넘겨 지나침(=건너뜀)
                bucket_key = f"{today_str}T{bm//60:02d}:{bm%60:02d}"
                if rule.get('lastFiredBucket') == bucket_key:
                    continue  # 이미 이 시각엔 발송함
                _fire_rule(rule)
                rule['lastFiredBucket'] = bucket_key
                rule['lastFiredAt']     = now.isoformat()
                fired_any = True
        except Exception as e:
            print(f"[예약발송 규칙 처리 오류] {rule.get('title','')}: {e}")

    if fired_any:
        with _SCHEDULE_LOCK:
            save_schedule_rules(SCHEDULE_RULES)


def _scheduler_loop():
    while True:
        try:
            _scheduler_tick()
        except Exception as e:
            print(f"[예약발송 스케줄러 오류] {e}")
        time.sleep(60)


# ── 설정 암호화 (Google Drive 저장용) ────────────────────────
@app.route('/telegram/encrypt', methods=['POST'])
def telegram_encrypt():
    if not CRYPTO_OK:
        return jsonify({'ok': False, 'error': 'cryptography 미설치: pip install cryptography'}), 500
    data     = request.json or {}
    password = data.get('password', '')
    config   = data.get('config',   {})
    if not password:
        return jsonify({'ok': False, 'error': '비밀번호 필요'}), 400
    try:
        encrypted = encrypt_data(config, password)
        return jsonify({'ok': True, 'encrypted': encrypted})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500


# ── Telegram 설정 복호화 (Google Drive 로드용) ────────────────────────
@app.route('/telegram/decrypt', methods=['POST'])
def telegram_decrypt():
    global TELEGRAM_TOKEN, TELEGRAM_CHAT_ID, TELEGRAM_MEMBERS
    if not CRYPTO_OK:
        return jsonify({'ok': False, 'error': 'cryptography 미설치: pip install cryptography'}), 500
    data      = request.json or {}
    password  = data.get('password',  '')
    encrypted = data.get('encrypted', '')
    if not password or not encrypted:
        return jsonify({'ok': False, 'error': 'password, encrypted 필요'}), 400
    try:
        config           = decrypt_data(encrypted, password)
        TELEGRAM_TOKEN   = config.get('token',           '')
        TELEGRAM_CHAT_ID = config.get('default_chat_id', '')
        TELEGRAM_MEMBERS = config.get('members',         [])
        save_tg_config(config)
        return jsonify({'ok': True, 'message': f'Telegram 설정 로드 완료 ({len(TELEGRAM_MEMBERS)}명)',
                        'members_count': len(TELEGRAM_MEMBERS)})
    except Exception as e:
        err = str(e)
        if 'InvalidToken' in err or 'token' in err.lower():
            err = '비밀번호가 틀렸거나 파일이 손상되었습니다'
        return jsonify({'ok': False, 'error': err}), 400


# ── SMTP 설정 암호화 (Google Drive 저장용) ───────────────────
@app.route('/mail/encrypt', methods=['POST'])
def mail_encrypt():
    if not CRYPTO_OK:
        return jsonify({'ok': False, 'error': 'cryptography 미설치'}), 500
    data     = request.json or {}
    password = data.get('password', '')
    if not password:
        return jsonify({'ok': False, 'error': '비밀번호 필요'}), 400
    try:
        config    = {'host': SMTP_HOST, 'port': SMTP_PORT,
                     'user': SMTP_USER, 'pass': SMTP_PASS}
        encrypted = encrypt_data(config, password)
        return jsonify({'ok': True, 'encrypted': encrypted})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500


# ── SMTP 설정 복호화 (Google Drive 로드용) ───────────────────
@app.route('/mail/decrypt', methods=['POST'])
def mail_decrypt():
    global SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASS
    if not CRYPTO_OK:
        return jsonify({'ok': False, 'error': 'cryptography 미설치'}), 500
    data      = request.json or {}
    password  = data.get('password',  '')
    encrypted = data.get('encrypted', '')
    if not password or not encrypted:
        return jsonify({'ok': False, 'error': 'password, encrypted 필요'}), 400
    try:
        config    = decrypt_data(encrypted, password)
        SMTP_HOST = config.get('host', SMTP_HOST)
        SMTP_PORT = int(config.get('port', SMTP_PORT))
        SMTP_USER = config.get('user', SMTP_USER)
        SMTP_PASS = config.get('pass', SMTP_PASS)
        # 로컬 파일에도 저장
        with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
            json.dump({'host': SMTP_HOST, 'port': SMTP_PORT,
                       'user': SMTP_USER, 'pass': SMTP_PASS},
                      f, ensure_ascii=False)
        return jsonify({'ok': True, 'message': f'SMTP 설정 로드 완료 ({SMTP_USER})'})
    except Exception as e:
        err = str(e)
        if 'InvalidToken' in err or 'token' in err.lower():
            err = '비밀번호가 틀렸거나 파일이 손상되었습니다'
        return jsonify({'ok': False, 'error': err}), 400


# ── 전체 설정 암호화 (mail + telegram 한 번에) ───────────────
@app.route('/all/encrypt', methods=['POST'])
def all_encrypt():
    if not CRYPTO_OK:
        return jsonify({'ok': False, 'error': 'cryptography 미설치'}), 500
    data     = request.json or {}
    password = data.get('password', '')
    if not password:
        return jsonify({'ok': False, 'error': '비밀번호 필요'}), 400
    try:
        mail_cfg = {'host': SMTP_HOST, 'port': SMTP_PORT,
                    'user': SMTP_USER, 'pass': SMTP_PASS}
        tg_cfg   = {'token': TELEGRAM_TOKEN,
                    'default_chat_id': TELEGRAM_CHAT_ID,
                    'members': TELEGRAM_MEMBERS}
        return jsonify({
            'ok':               True,
            'mail_encrypted':   encrypt_data(mail_cfg, password),
            'tg_encrypted':     encrypt_data(tg_cfg,   password)
        })
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500


# ── 전체 설정 복호화 (mail + telegram 한 번에) ───────────────
@app.route('/all/decrypt', methods=['POST'])
def all_decrypt():
    global SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASS
    global TELEGRAM_TOKEN, TELEGRAM_CHAT_ID, TELEGRAM_MEMBERS
    if not CRYPTO_OK:
        return jsonify({'ok': False, 'error': 'cryptography 미설치'}), 500
    data           = request.json or {}
    password       = data.get('password',       '')
    mail_encrypted = data.get('mail_encrypted', '')
    tg_encrypted   = data.get('tg_encrypted',   '')
    if not password:
        return jsonify({'ok': False, 'error': '비밀번호 필요'}), 400

    results = {}

    # SMTP 복호화
    if mail_encrypted:
        try:
            mc        = decrypt_data(mail_encrypted, password)
            SMTP_HOST = mc.get('host', SMTP_HOST)
            SMTP_PORT = int(mc.get('port', SMTP_PORT))
            SMTP_USER = mc.get('user', SMTP_USER)
            SMTP_PASS = mc.get('pass', SMTP_PASS)
            with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
                json.dump({'host': SMTP_HOST, 'port': SMTP_PORT,
                           'user': SMTP_USER, 'pass': SMTP_PASS},
                          f, ensure_ascii=False)
            results['mail'] = {'ok': True, 'message': f'SMTP 로드 완료 ({SMTP_USER})'}
        except Exception as e:
            results['mail'] = {'ok': False, 'error': str(e)}

    # Telegram 복호화
    if tg_encrypted:
        try:
            tc               = decrypt_data(tg_encrypted, password)
            TELEGRAM_TOKEN   = tc.get('token',           '')
            TELEGRAM_CHAT_ID = tc.get('default_chat_id', '')
            TELEGRAM_MEMBERS = tc.get('members',         [])
            save_tg_config(tc)
            results['telegram'] = {'ok': True,
                                   'message': f'Telegram 로드 완료 ({len(TELEGRAM_MEMBERS)}명)'}
        except Exception as e:
            results['telegram'] = {'ok': False, 'error': str(e)}

    overall_ok = all(v.get('ok') for v in results.values())
    return jsonify({'ok': overall_ok, 'results': results})


# ── SAP GUI Scripting 연동 (AI 문답 "SAP" 조회) ──────────────────
#    [2026-09-14 신규, 같은 날 32비트 브릿지로 재설계] AI 문답(js/04g~04h)에서 질문에 "SAP"가
#    언급되면 프런트가 이 엔드포인트를 호출한다. 비밀번호를 전혀 저장/입력하지 않는다 —
#    사용자가 평소처럼 SAP GUI에 직접 로그인해둔 "이미 열려 있는 세션"에 SAP GUI
#    Scripting(COM)으로 올라타서 지금 보이는 화면을 읽기만 한다.
#
#    ⚠️ 실제 COM 작업은 이 파일(64비트 Python)이 직접 하지 않는다 — sap_bridge_32.py를
#    32비트 Python 서브프로세스로 띄워서 시킨다. 이유(실사용 진단으로 확정): SAP GUI
#    Scripting의 COM 컴포넌트가 32비트로만 등록돼 있어서(레지스트리
#    HKLM\SOFTWARE\WOW6432Node\SAP\SAPGUI Front\...\Security 에만 UserScripting 값이 있고
#    64비트 레지스트리 뷰에는 없음), 이 백엔드처럼 64비트 Python에서 직접
#    win32com.client.GetObject("SAPGUI")를 부르면 항상 COM 오류 -2147221020(MK_E_SYNTAX)로
#    실패한다 — SAP GUI가 켜져 있고 로그인돼 있고 스크립팅도 켜져 있어도 100% 재현되는
#    비트수 불일치 문제였다(32비트 Python으로 같은 호출을 하면 즉시 성공 확인됨). 앞으로 이
#    엔드포인트를 고칠 때 실제 SAP 화면 파싱 로직(그리드/필드 덤프)은 전부
#    sap_bridge_32.py에 있으니 그쪽을 수정할 것 — 여기(kortek_backend.py)엔 절대 win32com을
#    다시 끌어오지 말 것(64비트라 어차피 못 씀).
_SAP_BRIDGE_PATH = os.path.join(BASE_DIR, 'sap_bridge_32.py')

# ── 승인원 표지 생성 — 자재그룹표(matgroups.json, 618건) ──────────────
#    [2026-09-15 신규] 사내에서 쓰던 별도 데스크톱 앱("연구소 가이드 시스템")의 exe를 분석해
#    그대로 재현한 기능 — MM03 화면엔 자재그룹 "코드"만 있고 명칭이 없어서, 그 앱에 내장돼
#    있던 618건짜리 자재그룹표(코드→대분류/중분류/소분류/설명)에서 코드로 찾는다. 원본 앱의
#    `_approval_group_name`과 동일하게, 표에 없는 코드(예: 사용금지된 코드)는 빈 값을
#    반환하고 화면/응답에 "확인 필요"로 표시해서 사람이 직접 판단하게 한다 — Item Description
#    앞부분으로 되짚는 등의 추측 로직은 원본 앱에서도 엉뚱한 이름이 채워지는 문제로 빠졌던
#    것이라 여기서도 넣지 않는다.
_MATGROUPS_PATH = os.path.join(BASE_DIR, 'matgroups.json')
_MATGROUPS_BY_CODE = {}
try:
    with open(_MATGROUPS_PATH, encoding='utf-8') as _mg_f:
        for _mg_entry in json.load(_mg_f):
            _code = (_mg_entry.get('code') or '').strip()
            if _code:
                _MATGROUPS_BY_CODE[_code] = _mg_entry
except Exception as _mg_err:
    print(f"[승인원 표지] matgroups.json 로드 실패 — 자재그룹명 조회가 안 됩니다: {_mg_err}")


def _approval_group_name(matkl):
    """자재그룹 코드(matkl)로 618건 표에서 소분류명(Item Group)을 찾는다. 못 찾으면 빈 문자열."""
    entry = _MATGROUPS_BY_CODE.get((matkl or '').strip())
    return (entry.get('sobun') or '') if entry else ''


# ══════════════════════════════════════════════════════════════
# 🆕 [2026-09-21, Phase 10] SAP/구매오더/AI 문답 이슈 수집 — 로컬 우편함(issue_events.jsonl)
#    설계: docs/phase10-issue-learning-design.md
#    - /sap-*, /po-* 호출 중 "실패·느림·레이아웃 미적용"만 한 줄씩 기록(성공은 기록 안 함).
#    - 백엔드는 Drive 토큰이 없으므로 여기에 쌓아두고, 브라우저가 /issue-drain으로 가져가
#      Drive에 올린 뒤 /issue-ack로 확정한다(최소 1회 전달 + 이벤트 id로 중복 제거).
#    - 저장 전 마스킹(자재번호·사업자번호·메일·사용자 경로). 수집 코드는 어떤 경우에도 본 기능을
#      방해하지 않는다(모든 구간 try/except).
# ══════════════════════════════════════════════════════════════
_ISSUE_FILE        = os.path.join(BASE_DIR, 'issue_events.jsonl')
_ISSUE_CURSOR_FILE = os.path.join(BASE_DIR, 'issue_events.cursor')
_ISSUE_EXPORT_DIR  = r'C:\SAP_DMS\SAP이슈'
_ISSUE_MAX_BYTES   = 5 * 1024 * 1024
_ISSUE_SLOW_MS     = 60000
_ISSUE_LOCK        = threading.Lock()
_ISSUE_ACTIVE      = {}   # rid -> {'overlap': bool}  (SAP 호출이 동시에 겹쳤는지 — 세션 충돌 의심 표시용)
_ISSUE_SKIP_MASK_KEYS = {'id', 'ts', 'rid', 'route', 'kind', 'domain', 'v', 'stage', 'tcode', 'program', 'screen', 'user'}


def _issue_file_hash(name):
    try:
        with open(os.path.join(BASE_DIR, name), 'rb') as f:
            return hashlib.sha1(f.read().replace(b'\r\n', b'\n')).hexdigest()[:8]
    except Exception:
        return ''


_ISSUE_ENV = {'backend': _issue_file_hash('kortek_backend.py'), 'bridge': _issue_file_hash('sap_bridge_32.py')}


def _issue_mask(text, limit=400):
    """저장 전 마스킹 — 사업자번호/5자리↑ 숫자열/메일/사용자 경로."""
    try:
        t = str(text)
        t = re.sub(r'\b\d{3}-\d{2}-\d{5}\b', '#biz', t)  # 하이픈 형태만 — 하이픈 없는 10자리(COM 오류코드 등)는 아래 5자리↑ 규칙이 '#'로 처리
        t = re.sub(r'\d{5,}', '#', t)
        t = re.sub(r'[\w.+-]+@[\w-]+\.[\w.-]+', '#mail', t)
        t = re.sub(r'(?i)([A-Z]:[\\/]Users[\\/])[^\\/\s"\']+', r'\1<user>', t)
        return t[:limit]
    except Exception:
        return ''


def _issue_norm(text):
    """군집 시그니처용 정규화 — 마스킹 + 따옴표 안 문자열/공백 정리."""
    t = _issue_mask(text, 300)
    t = re.sub(r'"[^"]*"', '"?"', t)
    t = re.sub(r"'[^']*'", "'?'", t)
    return re.sub(r'\s+', ' ', t).strip()[:120]


def _issue_mask_deep(obj, key=None):
    if isinstance(obj, str):
        return obj if key in _ISSUE_SKIP_MASK_KEYS else _issue_mask(obj)
    if isinstance(obj, dict):
        return {k: _issue_mask_deep(v, k) for k, v in obj.items()}
    if isinstance(obj, list):
        return [_issue_mask_deep(v, key) for v in obj]
    return obj


def _issue_params():
    """요청 파라미터의 "형태"만(값이 아니라 개수/열거형) — 기밀 값은 저장하지 않는다."""
    p = {}
    try:
        for k in ('material', 'materials'):
            v = request.args.get(k)
            if v:
                p['materialCount'] = len([x for x in re.split(r'[,\s]+', v) if x])
        for k in ('tcode', 'explosion', 'show_price', 'show_location', 'type', 'loc', 'plant', 'layout'):
            v = request.args.get(k)
            if v is not None:
                p[k] = v[:40]
        pat = request.args.get('pattern')
        if pat:
            p['patternLen'] = len(pat)
            p['patternHasWildcard'] = '*' in pat
        if request.is_json:
            j = request.get_json(silent=True) or {}
            if isinstance(j.get('items'), list):
                p['itemCount'] = len(j['items'])
            if isinstance(j.get('rows'), list):
                p['rowCount'] = len(j['rows'])
            for k in ('currency', 'plant', 'purchasingOrg'):
                if j.get(k):
                    p[k] = str(j[k])[:20]
    except Exception:
        pass
    return p


def _issue_append(event):
    """이벤트 1건을 우편함에 추가. 5MB 초과 시 오래된 절반을 버린다."""
    try:
        line = json.dumps(_issue_mask_deep(event), ensure_ascii=False) + '\n'
        with _ISSUE_LOCK:
            try:
                if os.path.exists(_ISSUE_FILE) and os.path.getsize(_ISSUE_FILE) > _ISSUE_MAX_BYTES:
                    with open(_ISSUE_FILE, 'rb') as f:
                        lines = f.read().splitlines(True)
                    with open(_ISSUE_FILE, 'wb') as f:
                        f.writelines(lines[len(lines) // 2:])
                    with open(_ISSUE_CURSOR_FILE, 'w') as f:
                        f.write('0')  # 오프셋이 바뀌므로 커서 초기화(중복은 id로 제거됨)
            except Exception:
                pass
            with open(_ISSUE_FILE, 'ab') as f:
                f.write(line.encode('utf-8'))
    except Exception as e:
        print(f"[이슈 수집] 기록 실패(무시): {e}")


def _issue_read_cursor():
    try:
        with open(_ISSUE_CURSOR_FILE, 'r') as f:
            return max(0, int(f.read().strip() or '0'))
    except Exception:
        return 0


@app.before_request
def _issue_before_request():
    try:
        path = request.path
        if request.method == 'OPTIONS' or not (path.startswith('/sap-') or path.startswith('/po-')):
            return None
        g._issue_t0 = time.time()
        g._issue_rid = (request.args.get('_rid') or '')[:40] or ('b_' + uuid.uuid4().hex[:8])
        with _ISSUE_LOCK:
            overlap = bool(_ISSUE_ACTIVE)
            for st in _ISSUE_ACTIVE.values():
                st['overlap'] = True
            _ISSUE_ACTIVE[g._issue_rid] = {'overlap': overlap}
    except Exception:
        pass
    return None


@app.after_request
def _issue_after_request(resp):
    try:
        rid = getattr(g, '_issue_rid', None)
        if not rid:
            return resp
        with _ISSUE_LOCK:
            st = _ISSUE_ACTIVE.pop(rid, {'overlap': False})
        dur_ms = int((time.time() - getattr(g, '_issue_t0', time.time())) * 1000)
        data = resp.get_json(silent=True) if (resp.mimetype or '').endswith('json') else None
        ok = data.get('ok') if isinstance(data, dict) else None
        path = request.path
        if resp.status_code >= 400 or ok is False:
            kind = 'call_fail'
        elif isinstance(data, dict) and data.get('layoutApplied') is False:
            kind = 'degraded_layout'
        elif dur_ms > _ISSUE_SLOW_MS:
            kind = 'call_slow'
        else:
            return resp
        err = ''
        if isinstance(data, dict):
            err = str(data.get('error') or data.get('layoutDiag') or '')
        elif resp.status_code >= 400:
            err = f'HTTP {resp.status_code} (JSON 아님)'
        snap = getattr(g, 'issue_snap', None)
        event = {
            'id': 'ev_%s_%s' % (format(int(time.time() * 1000), 'x'), uuid.uuid4().hex[:4]),
            'v': 1,
            'ts': datetime.now(KST).isoformat(timespec='seconds'),
            'domain': 'po' if path.startswith('/po-') else 'sap',
            'kind': kind,
            'rid': rid,
            'route': path,
            'params': _issue_params(),
            'result': {'ok': ok, 'http': resp.status_code, 'durMs': dur_ms,
                       'errorRaw': _issue_mask(err), 'errorNorm': _issue_norm(err)},
            'env': dict(_ISSUE_ENV),
        }
        if isinstance(data, dict) and 'layoutApplied' in data:
            event['result']['layout'] = {'applied': data.get('layoutApplied'), 'requested': data.get('layoutRequested'), 'diag': data.get('layoutDiag')}
        if getattr(g, 'issue_stage', None):
            event['stage'] = g.issue_stage
        if snap:
            event['snapshot'] = snap
        if st.get('overlap'):
            event['flags'] = ['concurrent_session_suspect']
        _issue_append(event)
    except Exception as e:
        print(f"[이슈 수집] after_request 실패(무시): {e}")
    return resp


@app.route('/issue-drain', methods=['GET'])
def issue_drain():
    """브라우저가 미전달 이벤트를 가져간다. 응답의 next를 /issue-ack로 돌려줘야 확정된다."""
    try:
        mx = max(1, min(int(request.args.get('max', '200')), 500))
        cursor = _issue_read_cursor()
        events, consumed = [], cursor
        with _ISSUE_LOCK:
            if not os.path.exists(_ISSUE_FILE):
                return jsonify({'ok': True, 'events': [], 'next': 0})
            size = os.path.getsize(_ISSUE_FILE)
            if cursor > size:
                cursor = consumed = 0
            with open(_ISSUE_FILE, 'rb') as f:
                f.seek(cursor)
                while len(events) < mx:
                    raw = f.readline()
                    if not raw:
                        break
                    consumed += len(raw)
                    try:
                        events.append(json.loads(raw.decode('utf-8')))
                    except Exception:
                        continue
        return jsonify({'ok': True, 'events': events, 'next': consumed})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500


@app.route('/issue-ack', methods=['POST'])
def issue_ack():
    """업로드가 끝난 이벤트까지 커서를 확정. 전부 소진했으면 우편함 파일을 비운다."""
    try:
        cur = int((request.get_json(silent=True) or {}).get('cursor', 0))
        with _ISSUE_LOCK:
            size = os.path.getsize(_ISSUE_FILE) if os.path.exists(_ISSUE_FILE) else 0
            cur = max(0, min(cur, size))
            if size and cur >= size:
                open(_ISSUE_FILE, 'wb').close()
                cur = 0
            with open(_ISSUE_CURSOR_FILE, 'w') as f:
                f.write(str(cur))
        return jsonify({'ok': True, 'cursor': cur})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500


@app.route('/issue-export', methods=['POST'])
def issue_export():
    """이슈 리포트(군집 요약 JSON)를 C:\\SAP_DMS\\SAP이슈\\ 에 저장 — 브라우저 다운로드가 아니라
    서버 저장으로 다른 SAP 기능들과 저장 위치를 통일한다."""
    try:
        body = request.get_json(silent=True) or {}
        data = body.get('data')
        if data is None:
            return jsonify({'ok': False, 'error': 'data가 없습니다.'}), 400
        name = re.sub(r'[^0-9A-Za-z_.\-가-힣]', '_', str(body.get('fileName') or ''))[:80]
        if not name.lower().endswith('.json'):
            name = (name or 'digest_' + datetime.now(KST).strftime('%Y%m%d')) + '.json'
        os.makedirs(_ISSUE_EXPORT_DIR, exist_ok=True)
        path = os.path.join(_ISSUE_EXPORT_DIR, name)
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=1)
        try:
            os.startfile(_ISSUE_EXPORT_DIR)
        except Exception:
            pass
        return jsonify({'ok': True, 'path': path})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 500


def _run_sap_bridge(extra_args, timeout, log_prefix):
    """sap_bridge_32.py를 32비트 Python 서브프로세스로 실행하고 JSON 결과를 돌려주는 공용
    헬퍼 — /sap-fetch와 /sap-open-document가 똑같이 쓴다(2026-09-14, 두 번째 엔드포인트
    추가하면서 중복 제거)."""
    try:
        proc = subprocess.run(
            ['py', '-3-32', _SAP_BRIDGE_PATH] + extra_args,
            capture_output=True, text=True, timeout=timeout, encoding='utf-8', errors='replace'
        )
    except FileNotFoundError:
        return {'ok': False, 'error': 'Python 런처(py.exe)를 찾을 수 없습니다 — Python 공식 설치 상태를 확인하세요.'}, 500
    except subprocess.TimeoutExpired:
        return {'ok': False, 'error': f'{timeout}초 안에 끝나지 않았습니다. SAP GUI에 응답 대기 중인 팝업이 떠 있지 않은지 확인해주세요.'}, 500

    stdout_lines = [ln for ln in (proc.stdout or '').strip().splitlines() if ln.strip()]
    if not stdout_lines:
        # 💡 32비트 Python/pywin32가 아예 없는 경우 등 — sap_bridge_32.py가 뭘 출력하기도 전에
        #    런처 자체가 실패한 상황. kortek_backend.bat의 자동 설치 블록을 다시 타라고 안내.
        stderr_tail = (proc.stderr or '').strip()
        print(f"[{log_prefix} 실패] 서브프로세스가 출력 없이 종료(returncode={proc.returncode}). stderr: {stderr_tail}")
        if 'No runtime installed' in stderr_tail or 'not installed' in stderr_tail.lower():
            return {'ok': False, 'error': '32비트 Python이 이 PC에 설치되어 있지 않습니다. kortek_backend.bat을 다시 실행하면 자동 설치됩니다(최초 1회, 몇 분 소요).'}, 500
        return {'ok': False, 'error': f'{log_prefix} 서브프로세스 실행에 실패했습니다. 백엔드 콘솔 창의 [{log_prefix} 실패] 로그를 확인하세요. (' + (stderr_tail[-300:] if stderr_tail else '상세 정보 없음') + ')'}, 500

    try:
        data = json.loads(stdout_lines[-1])
    except Exception:
        print(f"[{log_prefix} 실패] sap_bridge_32.py 출력 JSON 파싱 실패. stdout: {proc.stdout!r} stderr: {proc.stderr!r}")
        return {'ok': False, 'error': f'{log_prefix} 결과를 해석하지 못했습니다. 백엔드 콘솔 창의 로그를 확인하세요.'}, 500

    try:  # Phase 10: 브라우저 응답에는 싣지 않고 이슈 이벤트 로그로만 보낸다
        if isinstance(data, dict):
            g.issue_snap = data.pop('snapshot', None)
            g.issue_stage = data.pop('stage', None)
    except Exception:
        pass
    if data.get('ok'):
        print(f"[{log_prefix}] {data}")
    else:
        print(f"[{log_prefix} 실패] {data.get('error')}")
    return data, 200


@app.route('/sap-fetch', methods=['GET'])
def sap_fetch():
    data, status = _run_sap_bridge([], 30, 'SAP 조회')
    return jsonify(data), status


@app.route('/sap-open-document', methods=['GET'])
def sap_open_document():
    # 💡 [2026-09-14 신규] "SAP에서 P01 문서 열어줘" — MM03에서 이미 열어둔 자재의 "문서 데이터"
    #    탭에서 지정한 문서 타입(예: P01/C04/Q11)을 찾아 열고, 첨부된 원본 파일을 더블클릭해서
    #    연결된 프로그램(Acrobat 등)으로 바로 연다. 실제 클릭 순서는 SAP GUI "기록 및 재생"으로
    #    녹화한 매크로를 일반화해서 sap_bridge_32.py의 open_document()에 구현돼 있다.
    # 💡 [2026-09-15 확장] material 파라미터를 같이 주면(예: "SAP에서 106188 P01 문서 열어줘")
    #    사람이 화면을 미리 열어둘 필요 없이 sap_bridge_32.py가 직접 MM03으로 이동해 그 자재를
    #    조회한다(_navigate_to_material_document_tab, 2026-09-15 실사용 화면 녹화로 얻은
    #    컨트롤을 일반화). material 없이 호출하면 기존처럼 "이미 열려 있는 화면"을 그대로 쓴다.
    doc_type = (request.args.get('type') or '').strip()
    material = (request.args.get('material') or '').strip()
    if not doc_type:
        return jsonify({'ok': False, 'error': '문서 타입(type 파라미터)이 필요합니다. 예: /sap-open-document?type=P01'}), 400
    extra_args = ['open_document', doc_type]
    if material:
        extra_args.append(material)
    # 자재번호부터 직접 조회하는 경로는 화면 전환이 하나 더 있어 기존보다 시간이 더 걸릴 수
    # 있어 타임아웃을 넉넉히 둔다.
    timeout = 60 if material else 45
    data, status = _run_sap_bridge(extra_args, timeout, 'SAP 문서 열기')
    return jsonify(data), status


@app.route('/sap-material-documents', methods=['GET'])
def sap_material_documents():
    # 💡 [2026-09-15 신규] "SAP에서 106188 품번 정보 및 파일 열어줘"처럼 문서 타입(P01 등)을
    #    모른 채 자재번호만 말했을 때 — /sap-open-document처럼 특정 문서를 바로 열지 않고,
    #    먼저 "문서 데이터" 탭에 어떤 문서들이 있는지 목록만 읽어와 보여준다. 사람이 그 목록을
    #    보고 원하는 타입을 골라 "OO 문서 열어줘"라고 다시 말하면 /sap-open-document로 연다.
    material = (request.args.get('material') or '').strip()
    if not material:
        return jsonify({'ok': False, 'error': '자재번호(material 파라미터)가 필요합니다. 예: /sap-material-documents?material=106188'}), 400
    data, status = _run_sap_bridge(['fetch_material_documents', material], 30, 'SAP 자재 문서 목록 조회')
    return jsonify(data), status


@app.route('/sap-download-documents-batch', methods=['GET'])
def sap_download_documents_batch():
    # 💡 [2026-09-15 신규] "SAP에서 133012, 133010, 101831 문서 다운로드해줘"처럼 자재번호를
    #    여러 개 말했을 때 — MM03을 자재마다 드릴다운하는 대신 회사 SAP의 전용 배치 리포트
    #    ZDMSR004("DMS 첨부파일 일괄 다운로드 프로그램")를 실행해 한 번에 C:\SAP_DMS\로
    #    다운로드한다. 실제 컨트롤 조작은 sap_bridge_32.py의 download_documents_batch()에
    #    있다(2026-09-15 실사용 SAP GUI "기록 및 재생" 매크로로 확보한 정확한 ID 재현).
    materials_str = (request.args.get('materials') or '').strip()
    doc_type = (request.args.get('type') or 'P01').strip()
    if not materials_str:
        return jsonify({'ok': False, 'error': '자재번호 목록(materials 파라미터, 쉼표구분)이 필요합니다. 예: /sap-download-documents-batch?materials=133012,133010&type=P01'}), 400
    materials = [m.strip() for m in materials_str.split(',') if m.strip()]
    # 자재 수에 비례해 타임아웃을 넉넉히 둔다(팝업 입력+스크롤+다운로드까지 포함).
    timeout = min(300, 40 + 8 * len(materials))
    data, status = _run_sap_bridge(['download_documents_batch', doc_type, ','.join(materials)], timeout, 'SAP 문서 일괄 다운로드')
    return jsonify(data), status


@app.route('/sap-download-documents-by-pattern', methods=['GET'])
def sap_download_documents_by_pattern():
    # 💡 [2026-09-16 신규] "*01+01*500*로 조회된 아이템 승인원 다운로드해줘"처럼 자재번호를
    #    직접 나열하는 대신 자재내역(MAKT-MAKTX) 와일드카드 패턴으로 매치되는 자재를 전부
    #    찾아 문서를 다운로드한다. 자재번호를 구하는 단계(MM60의 자재번호 필드에서 F4로
    #    연 검색도움말을 라벨 매트릭스로 읽는 방식 — 실사용 라이브 진단으로 확인함, 위
    #    sap_bridge_32.py의 "자재내역 와일드카드 패턴으로..." 절 참고)와 다운로드 단계
    #    (기존 download_documents_batch 재사용, 7개씩 묶어 순차 실행)를 하나로 묶은
    #    download_documents_by_pattern()을 그대로 호출한다.
    pattern = (request.args.get('pattern') or '').strip()
    doc_type = (request.args.get('type') or 'P01').strip()
    if not pattern:
        return jsonify({'ok': False, 'error': '검색 패턴(pattern 파라미터)이 필요합니다. 예: /sap-download-documents-by-pattern?pattern=*01+01*500*&type=P01'}), 400
    # 검색(자재 수 미상, 최대 200건 안전장치) + 다운로드(7개씩 청크)까지 다 포함되므로
    # 넉넉하게 잡는다 — 실사용에서 매치 건수가 아주 많으면 이 값도 같이 늘려야 할 수 있음.
    data, status = _run_sap_bridge(['download_documents_by_pattern', pattern, doc_type], 240, 'SAP 패턴 검색 + 문서 일괄 다운로드')
    return jsonify(data), status


@app.route('/sap-bom', methods=['GET'])
def sap_bom():
    # 💡 [2026-09-15 신규] "SAP에서 502572 BOM 열어서 엑셀로 출력해줘"처럼 질문에 "BOM"과
    #    자재번호가 같이 언급되면(js/04h의 _aiFetchSapContext가 판정), 사람이 미리 SAP GUI에서
    #    그 자재의 BOM 화면을 열어둘 필요 없이 ZPP038("BOM 전개")로 직접 이동해 조회한다.
    #    실제 컨트롤 조작은 sap_bridge_32.py의 fetch_bom()/_navigate_to_bom_screen()에 있다
    #    (2026-09-15 실사용 SAP GUI "기록 및 재생" 매크로로 확보한 정확한 ID 재현).
    #    material 파라미터에 쉼표로 여러 자재를 같이 주면("SAP에서 502572,502573,502574 BOM
    #    보여줘"처럼 질문에 자재번호가 2개 이상이면) ZPP038의 자재코드 "복수 선택" 팝업으로
    #    한 번에 여러 BOM을 조회한다("BOM 복수 열람.vbs" 매크로로 확보, 2026-09-15).
    #    [2026-09-15 신규] tcode(single/multi/auto), explosion(single/multi), show_price,
    #    show_location 파라미터 추가 — 사용자가 실제 ZPP038 초기화면 캡처(Explosion type/
    #    Option을 빨간 박스로 표시)를 보여주며 "조회 전에 이 옵션들을 먼저 물어보고 선택한
    #    대로 조회해달라"고 요청 + "복수는 ZPP038, 단일은 ZPP033을 쓴다"고 알려줘서 추가함
    #    (js/04h의 BOM 옵션 사전질문 draft가 이 값들을 채워서 넘긴다).
    material = (request.args.get('material') or '').strip()
    plant = (request.args.get('plant') or '1000').strip()
    tcode_mode = (request.args.get('tcode') or 'auto').strip()  # 'single'|'multi'|'auto'
    explosion = (request.args.get('explosion') or 'single').strip()  # 'single'|'multi'
    show_price = (request.args.get('show_price') or '0').strip()
    show_location = (request.args.get('show_location') or '0').strip()
    # 🆕 [2026-09-17 신규, 사용자 요청] ALV 레이아웃을 더 이상 /STD_MC로 하드코딩하지 않고
    # AI 문답의 BOM 옵션 드롭다운에서 사람이 고른 값을 받는다 — 안 오면(하위호환) 빈 문자열을
    # 그대로 넘겨 sap_bridge_32.py의 기본값(_BOM_LAYOUT_VARIANT)이 적용되게 함.
    layout = (request.args.get('layout') or '').strip()
    if not material:
        return jsonify({'ok': False, 'error': '자재번호(material 파라미터)가 필요합니다. 예: /sap-bom?material=502572'}), 400
    material_count = len([m for m in material.split(',') if m.strip()])
    timeout = 30 if material_count <= 1 else min(90, 30 + 10 * material_count)
    data, status = _run_sap_bridge(
        ['fetch_bom', material, plant, tcode_mode, explosion, show_price, show_location, layout],
        timeout, 'SAP BOM 조회')
    return jsonify(data), status


_SAP_EXPORT_OUT_DIR = os.path.join('C:\\SAP_DMS', 'SAP조회')


@app.route('/sap-save-export', methods=['POST'])
def sap_save_export():
    # 💡 [2026-09-16 신규, 사용자 요청 "SAP 관련 저장 경로는 C:\SAP_DMS로 통일해줘"] "엑셀로
    #    내보내줘" 로컬 명령(js/04h의 _exportSapDataToExcel)이 예전엔 브라우저의 XLSX.js로
    #    직접 다운로드를 트리거해서(서버가 저장 경로를 모름) 실제 파일이 브라우저 기본
    #    다운로드 폴더(대부분 %USERPROFILE%\Downloads)에 떨어졌다 — 승인원 표지/구매오더/
    #    ZDMSR004 배치 다운로드는 전부 C:\SAP_DMS\ 아래 자기 전용 폴더에 쓰는데 이 경로만
    #    유일하게 어긋나 있었다. 이제 프런트가 XLSX 파일을 base64로 인코딩해 여기로 보내면
    #    서버가 직접 C:\SAP_DMS\SAP조회\ 에 써서 나머지 SAP 기능들과 저장 위치를 통일한다 —
    #    "해당 폴더로 이동하시겠습니까?"라고 되묻던 예전 흐름(/open-downloads-folder)도
    #    더 이상 필요 없어져 제거함(경로를 이제 서버가 정확히 알므로 다른 SAP 기능들처럼
    #    바로 자동으로 열어준다).
    data = request.get_json(force=True) or {}
    file_name = (data.get('fileName') or '').strip()
    data_base64 = data.get('dataBase64') or ''
    if not file_name or not data_base64:
        return jsonify({'ok': False, 'error': 'fileName/dataBase64가 필요합니다.'}), 400
    safe_name = re.sub(r'[\\/:*?"<>|]', '_', file_name)
    try:
        os.makedirs(_SAP_EXPORT_OUT_DIR, exist_ok=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': f'저장 폴더({_SAP_EXPORT_OUT_DIR})를 만들지 못했습니다: {e}'}), 500
    out_path = os.path.join(_SAP_EXPORT_OUT_DIR, safe_name)
    try:
        with open(out_path, 'wb') as f:
            f.write(base64.b64decode(data_base64))
    except Exception as e:
        return jsonify({'ok': False, 'error': f'파일 저장 중 오류가 발생했습니다: {e}'}), 500
    try:
        os.startfile(_SAP_EXPORT_OUT_DIR)
    except Exception:
        pass
    return jsonify({'ok': True, 'path': out_path, 'folder': _SAP_EXPORT_OUT_DIR,
                    'message': f'"{safe_name}" 파일을 {_SAP_EXPORT_OUT_DIR} 폴더에 저장했습니다. 탐색기로 그 폴더를 열었습니다.'})


@app.route('/sap-where-used', methods=['GET'])
def sap_where_used():
    # 💡 [2026-09-15 신규] "SAP에서 303410 역전개/사용처 보여줘"처럼 질문에 "역전개"/"사용처"와
    #    자재번호가 같이 언급되면(js/04h의 _aiFetchSapContext가 판정), CS15("단일레벨
    #    사용처리스트")로 직접 이동해 그 자재가 어느 상위 품목에 쓰이는지 조회한다 — ZPP038
    #    BOM 정전개의 반대 방향. 실제 컨트롤 조작은 sap_bridge_32.py의
    #    fetch_where_used()/_navigate_to_where_used_screen()에 있다(2026-09-15 실사용 SAP
    #    GUI "기록 및 재생" 매크로로 확보 — 단, 자재번호 입력 필드는 그 매크로에 안 나와서
    #    추측한 것이니 실패하면 그 함수의 주석부터 확인할 것).
    # 💡 [2026-09-15 다중 자재 지원] material 파라미터에 쉼표로 여러 자재를 같이 주면
    #    fetch_where_used가 자재마다 CS15를 따로 실행해 이어붙인다(ZPP038 BOM처럼 SAP 표준
    #    "복수 선택" 팝업으로 한 번에 묶지 않는 이유는 sap_bridge_32.py의 fetch_where_used
    #    docstring 참고 — ZPP046 배치 리포트를 시도했으나 _sap_dump_grid의 500행 캡이 여러
    #    자재 결과에 걸쳐 공유되는 바람에 사용처가 많은 자재 하나가 캡을 다 차지해 나머지
    #    자재 결과가 잘려나가는 문제가 실사용에서 확인돼, 자재별로 따로 실행하는 방식을 택함).
    material = (request.args.get('material') or '').strip()
    plant = (request.args.get('plant') or '1000').strip()
    if not material:
        return jsonify({'ok': False, 'error': '자재번호(material 파라미터)가 필요합니다. 예: /sap-where-used?material=303410'}), 400
    material_count = len([m for m in material.split(',') if m.strip()])
    timeout = 30 if material_count <= 1 else min(150, 30 + 20 * material_count)
    data, status = _run_sap_bridge(['fetch_where_used', material, plant], timeout, 'SAP 사용처 조회')
    return jsonify(data), status


@app.route('/sap-where-used-batch', methods=['GET'])
def sap_where_used_batch():
    # 💡 [2026-09-15 신규, 사용자 요청] "다중 사용처 조회해줘"/"일괄 역전개"처럼 "역전개"/
    #    "사용처"를 다중/일괄/복수의 의미로 말하면(js/04h의 판정) ZPP046("자재 사용처
    #    일괄조회")으로 SAP 서버 쪽에서 한 번에 묶어 빠르게 조회한다 — 위 /sap-where-used
    #    (자재별 CS15 순차 실행, 기본값)보다 빠르지만 _sap_dump_grid의 500행 캡을 여러 자재
    #    결과 전체가 공유해서 사용처가 많은 자재가 있으면 나머지가 잘릴 수 있다(자세한 내용은
    #    sap_bridge_32.py의 fetch_where_used_batch docstring 참고) — 사용자가 이
    #    트레이드오프를 알고 "다중/일괄/복수"라고 명시했을 때만 이 경로를 탄다(기본 동작
    #    아님). 복수 선택 팝업 입력 행이 8개까지만 확인돼 최대 8개로 자름.
    material = (request.args.get('material') or '').strip()
    plant = (request.args.get('plant') or '1000').strip()
    if not material:
        return jsonify({'ok': False, 'error': '자재번호(material 파라미터)가 필요합니다. 예: /sap-where-used-batch?material=303410,303411'}), 400
    materials = [m.strip() for m in material.split(',') if m.strip()][:8]
    data, status = _run_sap_bridge(['fetch_where_used_batch', ','.join(materials), plant], 60, 'SAP 사용처 일괄조회')
    return jsonify(data), status


@app.route('/sap-zmm009', methods=['GET'])
def sap_zmm009():
    # 💡 [2026-09-17 신규, 사용자 요청] "133025,133026 엑셀 출력해줘"처럼 자재번호(들)만
    #    주어지고 BOM/사용처/역전개 키워드가 없는 요청의 기본 조회 경로 — 사용자가 "ZMM009를
    #    기본으로 쓰면 좋겠다"고 요청해서 추가함. 실제 컨트롤 조작은 sap_bridge_32.py의
    #    fetch_zmm009_material_list()에 있다(사용자가 준 SAP GUI "기록 및 재생" 매크로에서
    #    그대로 가져온 정확한 필드 ID/버튼 시퀀스 — 추측 아님).
    material = (request.args.get('material') or '').strip()
    loc = (request.args.get('loc') or '1000').strip()
    if not material:
        return jsonify({'ok': False, 'error': '자재번호(material 파라미터)가 필요합니다. 예: /sap-zmm009?material=133025,133026'}), 400
    materials = [m.strip() for m in material.split(',') if m.strip()]
    timeout = min(150, 30 + 15 * len(materials))
    data, status = _run_sap_bridge(['fetch_zmm009_material_list', ','.join(materials), loc], timeout, 'SAP ZMM009 자재 조회')
    return jsonify(data), status


# ── 승인원 표지 생성 ──────────────────────────────────────────
#    [2026-09-15 신규] "SAP에서 104477 승인원 표지 생성해줘" — 사내 별도 데스크톱 앱("연구소
#    가이드 시스템")의 exe를 pyinstxtractor-ng로 풀고 main.pyc 바이트코드를 분석해서 그
#    앱의 GuiApp._approval_* 메서드들을 그대로 재현했다(디컴파일러가 Python 3.13을 지원 안
#    해서 완전한 소스 복원은 안 됐지만, docstring·상수 풀·바이트코드 흐름으로 충분히
#    재구성함 — 셀 주소·필드 ID·VKey 값까지 전부 원본에서 확인한 값 그대로). 두 단계로 나눔:
#    ① `/sap-approval-fetch` — SAP에서 자재정보만 조회(문서 생성 전 미리보기/확인용)
#    ② `/sap-approval-generate` — 조회된 정보 + 사람이 채운 나머지 항목으로 실제 엑셀/워드
#    파일을 만든다. AI 문답 쪽(js/04h)에서 이 둘을 연결해 "필요한 항목이 비어 있으면
#    물어보고, 다 채워지면 생성"하는 흐름을 만든다.
_APPROVAL_TEMPLATE_XLSX = os.path.join(BASE_DIR, 'templates', '승인원_양식.xlsx')
_APPROVAL_TEMPLATE_DOCX = os.path.join(BASE_DIR, 'templates', '승인원_양식.docx')
_APPROVAL_SHEET = '승인원표지'
_APPROVAL_OUT_DIR = os.path.join('C:\\SAP_DMS', '승인원표지')  # ZDMSR004 다운로드와 같은 자리(C:\SAP_DMS) 아래 전용 폴더
_APPROVAL_PRE_REMARK = '가승인원'
# 엑셀 셀 주소 — 원본 앱 바이트코드의 APPROVAL_CELLS 상수를 그대로 옮김.
_APPROVAL_CELLS = {
    'group_no': 'C8', 'group_name': 'C9', 'item_code': 'C10', 'desc': 'C11', 'sub_desc': 'C12',
    'dms': 'C13', 'remark': 'C14', 'chk_date': 'F20', 'chk_name': 'F21', 'apv_date': 'F28', 'apv_name': 'F29',
}
# 워드 표(중첩 표 0번, 1열=값) 행 번호 — 원본 앱의 APPROVAL_DOCX_ROWS.
_APPROVAL_DOCX_ROWS = {'group_no': 1, 'group_name': 2, 'item_code': 3, 'desc': 4, 'sub_desc': 5, 'dms': 6, 'remark': 7}
# 워드 표(중첩 표 1번) 서명란 (행, 열) — 원본 앱의 APPROVAL_DOCX_SIGN.
_APPROVAL_DOCX_SIGN = {'chk_date': (1, 3), 'chk_name': (2, 3), 'apv_date': (5, 3), 'apv_name': (6, 3)}


def _approval_fill_xlsx(tpl_path, out_path, r, ctx):
    """엑셀 양식에 값을 채워 저장한다(원본 앱 `_approval_fill_xlsx` 재현)."""
    import openpyxl
    wb = openpyxl.load_workbook(tpl_path)
    ws = wb[_APPROVAL_SHEET]
    ws[_APPROVAL_CELLS['item_code']] = r['code']
    ws[_APPROVAL_CELLS['desc']] = r.get('desc', '')
    ws[_APPROVAL_CELLS['sub_desc']] = r.get('sub', '')
    ws[_APPROVAL_CELLS['group_no']] = r.get('matkl', '')
    ws[_APPROVAL_CELLS['group_name']] = r.get('group', '')
    ws[_APPROVAL_CELLS['remark']] = ctx['remark']
    ws[_APPROVAL_CELLS['dms']] = ctx['dms_text']
    if ctx.get('name_chk'):
        ws[_APPROVAL_CELLS['chk_name']] = ctx['name_chk']
        ws[_APPROVAL_CELLS['chk_date']] = f"Date: {ctx['today']}"
    if ctx.get('name_apv'):
        ws[_APPROVAL_CELLS['apv_name']] = ctx['name_apv']
        ws[_APPROVAL_CELLS['apv_date']] = f"Date: {ctx['today']}"
    wb.save(out_path)


def _set_cell_text_preserve_format(cell, text):
    """`cell.text = 값` 대신 첫 문단의 첫 런(run)에만 텍스트를 채워 넣어 원본 서식(글꼴/크기/
    문단 간격 등)을 유지한다. ⚠️⚠️ [2026-09-15 실사용에서 확인] `cell.text = 값` 대입은
    python-docx 내부적으로 그 셀의 문단을 통째로 지우고 기본 서식의 새 문단/런을 만드는
    방식이라, 실제로 생성해보니 채워진 값들이 원본 라벨(맑은 고딕 12pt)과 다른 기본 폰트로
    나왔고 — 문단 간격 등도 같이 초기화되면서 누적된 높이 차이로 "1페이지로 나와야 하는
    양식이 2페이지로 밀려나는" 문제가 실사용에서 확인됨(원본 앱은 이 문제가 없었음, 즉 이
    코드의 재현 과정에서 생긴 회귀). 대신 기존 런이 있으면 그 런의 `.text`만 바꿔 서식을
    그대로 물려받고, 여분 런/문단은 제거해 항상 "문단 1개·런 1개"로 정리한다(런이 아예
    없던 빈 셀은 문단에 새 런을 추가 — 이 경우만 어쩔 수 없이 기본 서식이 적용됨)."""
    para = cell.paragraphs[0]
    for extra_para in cell.paragraphs[1:]:
        extra_para._element.getparent().remove(extra_para._element)
    if para.runs:
        para.runs[0].text = text
        for extra_run in para.runs[1:]:
            extra_run._element.getparent().remove(extra_run._element)
    else:
        para.add_run(text)


def _approval_fill_docx(tpl_path, out_path, r, ctx):
    """워드 양식에 값을 채워 저장한다(원본 앱 `_approval_fill_docx` 재현) — 양식은 큰 표 하나
    안에 중첩 표가 5개 들어있는 구조(python-docx로 실제 열어서 확인함): 0번이 본문(자재그룹/
    자재코드/... 7개 행), 1번이 서명란(담당자/팀장), 2~4번은 협력사 정보·문서번호·양식버전
    (이 코드에서는 건드리지 않음 — 협력사에서 별도로 채워 넣는 영역)."""
    from docx import Document
    doc = Document(tpl_path)
    outer_cell = doc.tables[0].cell(1, 0)
    nested = outer_cell.tables
    if len(nested) < 2:
        raise RuntimeError('워드 양식의 표 구조가 예상과 다릅니다.')
    t_main, t_sign = nested[0], nested[1]

    values = {
        'group_no': r.get('matkl', ''), 'group_name': r.get('group', ''), 'item_code': r['code'],
        'desc': r.get('desc', ''), 'sub_desc': r.get('sub', ''), 'dms': ctx['dms_text'], 'remark': ctx['remark'],
    }
    for key, row_idx in _APPROVAL_DOCX_ROWS.items():
        _set_cell_text_preserve_format(t_main.cell(row_idx, 1), str(values.get(key, '')))

    if ctx.get('name_chk'):
        rr, cc = _APPROVAL_DOCX_SIGN['chk_name']
        _set_cell_text_preserve_format(t_sign.cell(rr, cc), ctx['name_chk'])
        rr, cc = _APPROVAL_DOCX_SIGN['chk_date']
        _set_cell_text_preserve_format(t_sign.cell(rr, cc), f"Date: {ctx['today']}")
    if ctx.get('name_apv'):
        rr, cc = _APPROVAL_DOCX_SIGN['apv_name']
        _set_cell_text_preserve_format(t_sign.cell(rr, cc), ctx['name_apv'])
        rr, cc = _APPROVAL_DOCX_SIGN['apv_date']
        _set_cell_text_preserve_format(t_sign.cell(rr, cc), f"Date: {ctx['today']}")
    doc.save(out_path)


@app.route('/sap-approval-fetch', methods=['GET'])
def sap_approval_fetch():
    # "SAP에서 자재정보 가져오기" 단계 — 자재 여러 개(최대 30개, 원본 앱과 동일 상한)의
    # 자재내역·자재그룹·Sub-Description을 SAP에서 조회하고, 자재그룹명은 matgroups.json에서
    # 채운다. 파일 생성은 안 함 — /sap-approval-generate가 별도로 처리.
    materials_str = (request.args.get('materials') or '').strip()
    if not materials_str:
        return jsonify({'ok': False, 'error': '자재번호 목록(materials 파라미터, 쉼표구분)이 필요합니다.'}), 400
    materials = [m.strip() for m in materials_str.split(',') if m.strip()][:30]
    timeout = min(180, 20 + 8 * len(materials))
    data, status = _run_sap_bridge(['fetch_approval_info', ','.join(materials)], timeout, 'SAP 승인원 정보 조회')
    if data.get('ok') and isinstance(data.get('results'), list):
        for r in data['results']:
            if r.get('ok'):
                r['group'] = _approval_group_name(r.get('matkl', ''))
                r['group_missing'] = not bool(r['group'])
    return jsonify(data), status


@app.route('/sap-approval-generate', methods=['POST'])
def sap_approval_generate():
    # "표지 생성" 단계 — /sap-approval-fetch로 얻은 자재정보 결과(results)와 사람이 채운
    # 나머지 항목(출력형식/담당자/팀장/가승인원 여부/Revision/Remark)으로 실제 엑셀/워드
    # 파일을 만든다. 파일명 규칙은 원본 앱과 동일: (가)승인원_자재코드_Rev00.
    body = request.get_json(force=True, silent=True) or {}
    results = body.get('results') or []
    ok_results = [r for r in results if isinstance(r, dict) and r.get('ok') and r.get('code')]
    if not ok_results:
        return jsonify({'ok': False, 'error': '생성할 자재 정보가 없습니다. 먼저 SAP에서 자재정보를 조회해주세요.'}), 400

    fmt = (body.get('format') or 'both').strip().lower()
    if fmt not in ('xlsx', 'docx', 'both'):
        fmt = 'both'
    writer = (body.get('writer') or '').strip()
    leader = (body.get('leader') or '').strip()
    is_pre = bool(body.get('is_pre'))
    rev = (body.get('rev') or '00').strip() or '00'
    remark = (body.get('remark') or '').strip()
    kind = (body.get('kind') or '').strip()
    if kind not in ('New', 'Update'):
        kind = 'New' if rev in ('00', '0') else 'Update'  # Revision이 있으면 보통 Update로 보는 게 합리적

    # remark 가공 — 원본 앱 `_approval_generate` 로직 그대로(가승인원이면 그 문구를 보장, 완전히 비면 '.').
    if is_pre:
        if remark and _APPROVAL_PRE_REMARK not in remark:
            remark = f'{_APPROVAL_PRE_REMARK} / {remark}'
        elif not remark:
            remark = _APPROVAL_PRE_REMARK
    if not remark:
        remark = '.'

    dms_text = (('■' if kind == 'New' else '□') + ' New '
                + ('■' if kind == 'Update' else '□') + f' Update (Revision: {rev} )')
    ctx = {
        'today': datetime.now(KST).strftime('%Y.%m.%d'),
        'remark': remark, 'name_chk': writer, 'name_apv': leader, 'dms_text': dms_text,
    }

    need_xlsx = fmt in ('xlsx', 'both')
    need_docx = fmt in ('docx', 'both')
    if need_xlsx and not os.path.exists(_APPROVAL_TEMPLATE_XLSX):
        return jsonify({'ok': False, 'error': f'엑셀 양식({_APPROVAL_TEMPLATE_XLSX})을 찾을 수 없습니다.'}), 500
    if need_docx and not os.path.exists(_APPROVAL_TEMPLATE_DOCX):
        return jsonify({'ok': False, 'error': f'워드 양식({_APPROVAL_TEMPLATE_DOCX})을 찾을 수 없습니다.'}), 500

    prefix = _APPROVAL_PRE_REMARK if is_pre else '승인원'
    try:
        os.makedirs(_APPROVAL_OUT_DIR, exist_ok=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': f'저장 폴더({_APPROVAL_OUT_DIR})를 만들지 못했습니다: {e}'}), 500

    made, failed = [], []
    for r in ok_results:
        base = f"{prefix}_{r['code']}_Rev{rev}"
        if need_xlsx:
            out_path = os.path.join(_APPROVAL_OUT_DIR, base + '.xlsx')
            try:
                _approval_fill_xlsx(_APPROVAL_TEMPLATE_XLSX, out_path, r, ctx)
                made.append(out_path)
            except Exception as e:
                failed.append(f"{r['code']} (xlsx): {e}")
        if need_docx:
            out_path = os.path.join(_APPROVAL_OUT_DIR, base + '.docx')
            try:
                _approval_fill_docx(_APPROVAL_TEMPLATE_DOCX, out_path, r, ctx)
                made.append(out_path)
            except Exception as e:
                failed.append(f"{r['code']} (docx): {e}")

    # 다운로드 완료 후 폴더를 탐색기로 열어준다 — ZDMSR004 일괄 다운로드와 동일한 패턴.
    try:
        os.startfile(_APPROVAL_OUT_DIR)
    except Exception:
        pass

    msg = f'승인원 표지 {len(made)}건이 생성되었습니다.'
    if failed:
        msg += f'\n\n실패 {len(failed)}건:\n' + '\n'.join(failed[:5])
    return jsonify({'ok': True, 'made': made, 'failed': failed, 'folder': _APPROVAL_OUT_DIR, 'message': msg})


# 🛒 [2026-09-15 신규] "구매오더 요청" — AI 문답에 전자세금계산서/견적서 PDF를 첨부하면 항목을
#    추출해 "연구소 구매오더,기타출고 제안 BDC Upload양식"(사용자 제공 Z38MMR060.xls) 형식의
#    엑셀을 만들어준다. ZMMR060에 이 엑셀을 실제로 "업로드"하는 정확한 SAP GUI 조작은 아직
#    확인 전(매크로/라이브 진단 대기 중)이라, 이 엔드포인트는 **엑셀 생성 + 폴더 자동 오픈까지만**
#    한다 — SAP 자동화(업로드~F8~저장확인~ZMM018 발주서 출력)는 그 조작이 확인되면 이어서 붙일
#    예정(CLAUDE.md "🛒 구매오더 요청" 절 참고). 컬럼 순서/헤더 문구/시작 위치(헤더 5행·B열,
#    데이터 6행부터)는 사용자가 준 원본 Z38MMR060.xls의 "Upload" 시트를 xlrd로 직접 읽어
#    그대로 확인한 값 — 추측이 아니라 실측값이므로 바꾸지 말 것.
_PO_OUT_DIR = os.path.join('C:\\SAP_DMS', '구매오더')
_PO_EXCEL_HEADERS = ['자재코드', '자재명', '요청수량', '필요일자', '구매그룹', '프로젝트코드',
                      '수령인', '구매담당자 사번', '요청사유', 'VINA PO', '목적', '비고']


@app.route('/po-build-excel', methods=['POST'])
def po_build_excel():
    # 💡 원본 템플릿이 레거시 BIFF .xls라(openpyxl은 이 포맷을 못 씀) xlwt로 새로 써서 같은
    #    확장자(.xls)로 맞춘다 — SAP이 실제로 어떤 방식으로 이 파일을 읽을지 아직 몰라서, 형식
    #    불일치로 인한 새로운 실패를 피하려고 원본과 동일한 바이너리 포맷을 그대로 재현했다.
    import xlwt
    data = request.get_json(silent=True) or {}
    rows = data.get('rows') or []
    if not rows:
        return jsonify({'ok': False, 'error': '엑셀에 넣을 품목이 없습니다.'}), 400

    try:
        os.makedirs(_PO_OUT_DIR, exist_ok=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': f'저장 폴더({_PO_OUT_DIR})를 만들지 못했습니다: {e}'}), 500

    wb = xlwt.Workbook(encoding='utf-8')
    ws = wb.add_sheet('Upload')
    title_style = xlwt.easyxf('font: bold on')
    header_style = xlwt.easyxf('font: bold on; pattern: pattern solid, fore_colour yellow;')
    ws.write(1, 1, '연구소 구매오더,기타출고 제안 BDC Upload양식 [양식변경불가, 작성방법 참조]', title_style)
    ws.write(2, 1, '※ 양식변경사항 발생시, iCare로 요청 바랍니다.')
    for ci, h in enumerate(_PO_EXCEL_HEADERS):
        ws.write(4, 1 + ci, h, header_style)
    for ri, row in enumerate(rows):
        for ci, h in enumerate(_PO_EXCEL_HEADERS):
            ws.write(5 + ri, 1 + ci, row.get(h, ''))

    ts = datetime.now(KST).strftime('%Y%m%d_%H%M%S')
    file_name = f'구매오더요청_{ts}.xls'
    out_path = os.path.join(_PO_OUT_DIR, file_name)
    try:
        wb.save(out_path)
    except Exception as e:
        return jsonify({'ok': False, 'error': f'엑셀 저장 중 오류: {e}'}), 500

    # 생성 완료 후 폴더를 탐색기로 열어준다 — ZDMSR004/승인원 표지와 동일한 패턴.
    try:
        os.startfile(_PO_OUT_DIR)
    except Exception:
        pass

    return jsonify({'ok': True, 'fileName': file_name, 'path': out_path, 'folder': _PO_OUT_DIR,
                     'message': f'구매오더 요청 엑셀({len(rows)}개 품목)을 생성했습니다: {file_name}'})


@app.route('/po-sap-prepare', methods=['POST'])
def po_sap_prepare():
    # 🛒 [2026-09-15 신규] "구매오더 요청" 2단계 — /po-build-excel로 만든 엑셀을 ZMMR060에
    #    업로드해 협력사(사업자등록번호로 검색)/세금코드/단가까지 채우고 저장 직전에
    #    멈춘다(사용자와 합의한 "저장 직전 확인 후 정지" 설계 — CLAUDE.md "🛒 구매오더
    #    요청" 절 참고). 실제 저장은 /po-sap-confirm-save가 사람 확인 후 별도로 호출한다.
    data = request.get_json(silent=True) or {}
    excel_path = (data.get('excelPath') or '').strip()
    biz_reg_no = (data.get('bizRegNo') or '').strip()
    items = data.get('items') or []
    plant = (data.get('plant') or '1000').strip()
    currency = (data.get('currency') or 'KRW').strip() or 'KRW'
    if not excel_path or not biz_reg_no or not items:
        return jsonify({'ok': False, 'error': 'excelPath/bizRegNo/items가 모두 필요합니다.'}), 400
    timeout = min(120, 40 + 8 * len(items))
    data_out, status = _run_sap_bridge(
        ['prepare_po_from_excel', excel_path, biz_reg_no, json.dumps(items, ensure_ascii=False), plant, currency],
        timeout, 'SAP 구매오더 준비')
    return jsonify(data_out), status


@app.route('/po-sap-confirm-save', methods=['POST'])
def po_sap_confirm_save():
    # 🛒 [2026-09-15 신규] "구매오더 요청" 3단계 — 사람이 /po-sap-prepare 결과를 채팅에서
    #    확인하고 "저장해줘"라고 답한 뒤에만 호출된다. 실제 저장(구매오더 확정) + 오더번호
    #    확보 + ZMM018 발주서 PDF 출력까지 진행한다.
    data = request.get_json(silent=True) or {}
    purchasing_org = (data.get('purchasingOrg') or '9000').strip()
    plant = (data.get('plant') or '1000').strip()
    data_out, status = _run_sap_bridge(['confirm_save_po', purchasing_org, plant], 60, 'SAP 구매오더 저장')
    return jsonify(data_out), status


@app.route('/po-print-via-zmm018', methods=['POST'])
def po_print_via_zmm018():
    # 🛒 [2026-09-15 신규] 복구용 — 구매오더가 이미 저장됐는데(예: /po-sap-confirm-save가
    #    저장은 성공했지만 ZMM018 출력 단계에서만 실패한 경우, 또는 사람이 SAP에서 직접
    #    저장한 경우) 오더번호를 알고 있을 때 발주서 PDF 출력만 다시 시도한다. 저장(SAVE)이
    #    전혀 없는 순수 조회/출력 동작.
    data = request.get_json(silent=True) or {}
    po_number = (data.get('poNumber') or '').strip()
    purchasing_org = (data.get('purchasingOrg') or '9000').strip()
    plant = (data.get('plant') or '1000').strip()
    if not po_number:
        return jsonify({'ok': False, 'error': 'poNumber가 필요합니다.'}), 400
    data_out, status = _run_sap_bridge(['print_po_via_zmm018', po_number, purchasing_org, plant], 60, 'SAP 발주서 PDF 출력')
    return jsonify(data_out), status


# ── 백엔드 자동 업데이트("SAP 조회 연동" 절 kortek_backend.zip 배포 방식의 대안, 2026-09-15) ─
# 이 앱은 GitHub Pages(정적 프런트) + 각 PC의 로컬 백엔드(kortek_backend.py) 구조라, 백엔드
# 파일이 바뀔 때마다 사용자가 kortek_backend.zip을 다시 받아 기존 폴더에 덮어써야 했다 —
# 그 절차를 사람이 잊거나 귀찮아하면 다른 팀원 PC는 계속 구버전 백엔드로 남는다. 그래서 프런트가
# 구글 로그인 성공 직후(js/04b-core-app-drive-sync.js) 로컬 백엔드에게 "지금 GitHub main
# 브랜치의 최신 파일과 내용이 다른 게 있는지" 물어보고(`/self-check-update`), 다르면 배너를
# 띄워 사용자가 버튼을 누르면 로컬 백엔드가 스스로 그 파일들을 GitHub raw에서 받아 자기 자신의
# 디렉터리(BASE_DIR)에 덮어쓴다(`/self-update`) — 버전 번호를 별도로 관리하지 않고 매번 GitHub의
# 실제 파일 내용과 직접 바이트 비교하므로 "버전 올리는 걸 깜빡해서 갱신 감지가 안 되는" 문제가
# 없다.
# ⚠️⚠️ [2026-09-16 재검토 후 결정 변경] 원래는 "자동 재시작은 하지 않는다"(포트 점유 해제
# 타이밍 등 복잡도·위험도 대비 이득이 적다고 판단)였으나, SAP 디버깅 중 백엔드가 자주 바뀌고
# 다른 사용자 PC는 시작프로그램으로 한 번 켜진 뒤 사람이 재시작을 깜빡하기 쉽다는 실사용
# 문제(구버전 백엔드가 새 SAP 엔드포인트 요청에 HTML 404를 돌려줘 "Unexpected token JSON"
# 오류로 오인되던 사고)가 반복돼, 사용자 요청으로 자동 재시작을 다시 구현하기로 결정함 —
# 아래 `_spawn_restarted_backend()` 참고. 실패해도(새 프로세스 기동 실패 등) 예외 없이 조용히
# `restarting: False`로 응답해 프런트가 기존처럼 "수동으로 재시작해주세요" 안내로 자연히
# 폴백하게 한다.
_GITHUB_RAW_BASE = 'https://raw.githubusercontent.com/yhparkkortek/ganttchart/main'
# kortek_backend.zip에 포함되는 배포 대상 파일 목록과 동일하게 유지할 것(.claude/settings.json의
# PostToolUse 훅 Compress-Archive 목록 참고) — 새 배포 파일이 추가되면 여기도 같이 추가해야 함.
_SELF_UPDATE_FILES = [
    'kortek_backend.py',
    'sap_bridge_32.py',
    'kortek_backend.bat',
    'kortek_backend_install.bat',
    'kortek_backend_start_minimized.vbs',
    'matgroups.json',
    'templates/승인원_양식.xlsx',
    'templates/승인원_양식.docx',
]
# ⚠️⚠️ [2026-09-16 실사용 버그수정] Windows 배치(.bat)/VBScript(.vbs) 파일은 git이 텍스트
# 파일을 저장소 안에서는 항상 LF로 정규화해 저장하므로(로컬 작업 사본의 줄바꿈과 무관하게)
# GitHub raw도 항상 LF-only 바이트를 돌려준다 — 그걸 그대로 바이트 그대로 써버리면(기존 코드)
# 사용자 PC에는 LF-only .bat/.vbs 파일이 생긴다. cmd.exe의 배치 파서는 특히 이 파일들처럼
# 괄호로 감싼 `IF (...) ELSE (...)` 블록이 많을 때 LF-only 줄바꿈에서 오동작하거나
# ("...was unexpected at this time." 류 오류, 또는 블록이 조용히 실행 안 됨), 옛 스타일
# 메모장에서는 줄바꿈 자체를 인식 못해 텍스트가 한 줄로 뭉쳐 보인다("파일이 이상함"이라는
# 실사용 제보의 실제 원인으로 추정) — 이 확장자들은 항상 CRLF로 정규화해서 써야 한다.
_SELF_UPDATE_CRLF_EXTS = ('.bat', '.vbs')


def _fetch_github_raw(rel_path):
    url = _GITHUB_RAW_BASE + '/' + '/'.join(quote(seg) for seg in rel_path.split('/'))
    resp = requests.get(url, timeout=10)
    resp.raise_for_status()
    return resp.content


def _normalize_for_compare(rel_path, data):
    """`.bat`/`.vbs`는 로컬에 항상 CRLF로 써두는데(아래 `/self-update`의
    `_SELF_UPDATE_CRLF_EXTS` 참고) GitHub raw는 git이 LF로 정규화해 저장한 바이트를 그대로
    돌려주므로, 줄바꿈만 다르고 내용은 같은 경우까지 매번 "구버전"으로 오탐하지 않도록
    비교 직전에 양쪽 다 LF로 맞춰서 비교한다 — 실제 파일 갱신(`/self-update`)에는 영향
    없음(그쪽은 항상 CRLF로 쓰는 로직 그대로 유지)."""
    if rel_path.lower().endswith(_SELF_UPDATE_CRLF_EXTS):
        return data.replace(b'\r\n', b'\n')
    return data


@app.route('/self-check-update', methods=['GET'])
def self_check_update():
    changed, errors = [], []
    for rel_path in _SELF_UPDATE_FILES:
        try:
            remote_bytes = _fetch_github_raw(rel_path)
        except Exception as e:
            errors.append(f'{rel_path}: {e}')
            continue
        local_path = os.path.join(BASE_DIR, rel_path)
        local_bytes = b''
        if os.path.exists(local_path):
            with open(local_path, 'rb') as f:
                local_bytes = f.read()
        if _normalize_for_compare(rel_path, remote_bytes) != _normalize_for_compare(rel_path, local_bytes):
            changed.append(rel_path)
    return jsonify({'ok': True, 'outdated': len(changed) > 0, 'changedFiles': changed, 'errors': errors})


def _spawn_restarted_backend():
    """새 `kortek_backend.py` 프로세스를 별도 콘솔로 띄운다 — 새 프로세스는 포트 5000이
    아직 이전(지금 실행 중인) 프로세스에 점유돼 있으면 `_run_flask_with_port_retry`가
    자동으로 재시도하므로, 여기서는 타이밍을 정확히 맞출 필요 없이 그냥 곧바로 띄우기만
    하면 된다. `CREATE_NEW_CONSOLE`로 띄워서 사람이 재시작이 실제로 일어났음을(새 콘솔
    창이 뜸) 눈으로 확인할 수 있게 한다 — 완전히 숨기면(DETACHED_PROCESS) 조용히 실패해도
    아무도 모르는 게 더 위험하다고 판단."""
    script_path = os.path.join(BASE_DIR, 'kortek_backend.py')
    creationflags = 0
    if sys.platform == 'win32':
        creationflags = subprocess.CREATE_NEW_CONSOLE
    subprocess.Popen(
        [sys.executable, script_path],
        cwd=BASE_DIR,
        creationflags=creationflags,
        close_fds=True,
    )


def _delayed_self_exit(delay_sec=0.6):
    """지금 응답을 클라이언트에게 다 보낼 시간을 준 뒤(0.6초) 이 프로세스를 종료해 포트
    5000을 반납한다 — 응답을 먼저 안 보내고 바로 죽으면 요청자는 연결이 끊긴 것으로만
    보여 재시작이 실제로 성공했는지 알 방법이 없다."""
    def _exit():
        time.sleep(delay_sec)
        os._exit(0)
    threading.Thread(target=_exit, daemon=True).start()


@app.route('/self-update', methods=['POST'])
def self_update():
    updated, errors = [], []
    for rel_path in _SELF_UPDATE_FILES:
        try:
            remote_bytes = _fetch_github_raw(rel_path)
        except Exception as e:
            errors.append(f'{rel_path}: 다운로드 실패 - {e}')
            continue
        local_path = os.path.join(BASE_DIR, rel_path)
        try:
            os.makedirs(os.path.dirname(local_path), exist_ok=True)
            write_bytes = remote_bytes
            if rel_path.lower().endswith(_SELF_UPDATE_CRLF_EXTS):
                # 먼저 혹시 있을 CRLF를 LF로 되돌린 뒤 전부 CRLF로 바꿔서, 원본이 LF-only든
                # 이미 CRLF든 결과가 항상 깨끗한 CRLF가 되게 한다(CRLF를 또 CRLF화해서
                # CRCRLF가 되는 사고 방지).
                write_bytes = remote_bytes.replace(b'\r\n', b'\n').replace(b'\n', b'\r\n')
            with open(local_path, 'wb') as f:
                f.write(write_bytes)
            updated.append(rel_path)
        except Exception as e:
            errors.append(f'{rel_path}: 저장 실패 - {e}')

    ok = len(errors) == 0
    restarting = False
    if ok and updated:
        try:
            _spawn_restarted_backend()
            restarting = True
        except Exception as e:
            # 새 프로세스를 못 띄워도 파일 갱신 자체는 이미 끝났으니 오류로 취급하지 않고,
            # restarting:False만 알려서 프런트가 수동 재시작 안내로 폴백하게 한다.
            errors.append(f'자동 재시작 시도 실패(파일은 갱신됨, 수동 재시작 필요) - {e}')

    response = jsonify({'ok': ok, 'updatedFiles': updated, 'errors': errors, 'restarting': restarting})
    if restarting:
        # 응답을 다 만든 뒤에만 종료를 예약 — 아래 return으로 클라이언트에 응답이 나간 다음에
        # 이 스레드가 깨어나 프로세스를 종료한다.
        _delayed_self_exit()
    return response


def _run_flask_with_port_retry(max_wait_sec=15):
    """`/self-update`의 자동 재시작(`_spawn_restarted_backend`)이 새 프로세스를 띄운 시점에
    이전 프로세스가 아직 포트 5000을 반납하기 전일 수 있다(`_delayed_self_exit`이 0.6초 뒤
    종료하지만, 그 사이 새 프로세스가 먼저 뜰 수도 있음) — 정확한 타이밍을 맞추는 대신,
    바인드 실패(OSError, Windows에서는 통상 WinError 10048 "각 소켓 주소는 한 번만 사용할
    수 있습니다")를 잡아 포트가 열릴 때까지 짧게 재시도하는 쪽이 더 안전하다고 판단함."""
    start = time.time()
    while True:
        try:
            app.run(host='127.0.0.1', port=5000, debug=False)
            return
        except OSError as e:
            if time.time() - start > max_wait_sec:
                raise
            time.sleep(0.5)


# ══════════════════════════════════════════════════════════════
if __name__ == '__main__':
    print("=" * 58)
    print("  KORTEK Backend v1.0")
    print("  자동화 메일 · Telegram 알람 · AI 분석 서버")
    print("-" * 58)
    print(f"  SMTP : {SMTP_HOST or '(미설정)'}:{SMTP_PORT}")
    print(f"  계정 : {SMTP_USER or '(미설정)'}")
    print(f"  POP3 : {POP3_HOST}:{POP3_PORT}")
    print(f"  TG   : {'연결됨 ✓' if TELEGRAM_TOKEN else '미설정 (알람설정 > Telegram 탭에서 입력)'}")
    print(f"  팀원 : {len(TELEGRAM_MEMBERS)}명 등록")
    print(f"  예약 : {len(SCHEDULE_RULES)}건 등록 (1분마다 자동 확인)")
    print("-" * 58)
    print("  URL  : http://127.0.0.1:5000")
    print("  종료 : Ctrl+C 또는 창 닫기")
    print("=" * 58)
    threading.Thread(target=_scheduler_loop, daemon=True).start()
    _run_flask_with_port_retry()
